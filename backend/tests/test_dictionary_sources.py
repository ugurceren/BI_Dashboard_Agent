"""Sözlük kaynakları: rol başına birden çok tablo / sayfa, eksik kolonlar, normalize, Excel ve MySQL okuyucuları."""

import openpyxl
import pytest

from app.dictionary.sources import DEFAULTS, ExcelReader, MySQLReader, SqlServerReader, collect, default_sources, suggest_role


class FakeReader:
    """Tablo adı → satırlar (kolon adları küçük harf)."""

    kind = "sqlserver"
    valid = staticmethod(SqlServerReader.valid)

    def __init__(self, tables: dict[str, list[dict]]):
        self.tables = {k.lower(): v for k, v in tables.items()}
        self.cols = {k.lower(): set(v[0].keys()) if v else set() for k, v in tables.items()}

    def columns(self, names):
        return {n.lower(): self.cols[n.lower()] for n in names if n.lower() in self.cols}

    def read(self, name):
        return self.tables[name.lower()]

    def list_tables(self):
        return {k: sorted(v) for k, v in self.cols.items()}


def test_multiple_tables_merged_and_optional_columns_blank():
    r = FakeReader({
        "meta.dd_tables": [{"table_name": "dbo.FactSales", "business_name": "Satış", "row_count": "1.234"}],
        "satis.kolon": [{"table_name": "dbo.FactSales", "column_name": "Amount", "is_pii": "Hayır"}],
        "finans.kolon": [{"table_name": "dbo.FactSales", "column_name": "Email", "is_pii": "Evet"},
                         {"table_name": "", "column_name": "boş"}],
    })
    got = collect(r, {"tables": ["meta.dd_tables"], "columns": ["satis.kolon", "finans.kolon"], "relationships": [], "metrics": []})
    assert not got.errors
    assert got.counts == {"tables": 1, "columns": 2}
    t = got.rows["tables"][0]
    assert t["row_count"] == 1234 and t["description"] is None and t["table_type"] is None
    assert [c["is_pii"] for c in got.rows["columns"]] == [False, True]
    assert any("boş sayıldı" in w for w in got.warnings) and any("1 satır atlandı" in w for w in got.warnings)


def test_required_column_missing_and_unknown_table():
    r = FakeReader({"meta.x": [{"table_name": "a"}]})
    got = collect(r, {"tables": ["meta.x"], "columns": ["meta.x", "meta.yok"], "relationships": [], "metrics": []})
    assert any("'meta.x' tablosunda zorunlu kolon(lar) yok: column_name" in e for e in got.errors)
    assert any("'meta.yok' tablosu bulunamadı" in e for e in got.errors)
    assert any("Kolonlar için en az bir geçerli tablo" in e for e in got.errors)


def test_invalid_sql_names_rejected():
    r = FakeReader({})
    got = collect(r, {"tables": ["meta.t; DROP TABLE x"], "columns": [], "relationships": [], "metrics": []})
    assert any("Geçersiz ad" in e for e in got.errors)
    assert MySQLReader.valid("dd_tables") and MySQLReader.valid("meta.dd_tables") and not MySQLReader.valid("x`; DROP")


def test_excel_reader_with_turkish_sheets(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Tablolar"
    ws.append(["Table_Name", "Business Name", "Row_Count"])
    ws.append(["dbo.DimProduct", "Ürün", 606])
    ws2 = wb.create_sheet("Kolonlar")
    ws2.append([None])  # başlıktan önce boş satır
    ws2.append(["table_name", "column_name", "is_pii", "synonyms"])
    ws2.append(["dbo.DimProduct", "EnglishProductName", "", "ürün adı, isim"])
    ws2.append([None, None, None, None])
    ws3 = wb.create_sheet("İlişkiler")
    ws3.append(["from_table", "from_column", "to_table", "to_column", "is_active"])
    ws3.append(["dbo.FactSales", "ProductKey", "dbo.DimProduct", "ProductKey", "Hayır"])
    path = tmp_path / "sozluk.xlsx"
    wb.save(path)

    r = ExcelReader(path)
    assert set(r.list_tables()) == {"Tablolar", "Kolonlar", "İlişkiler"}
    got = collect(r, {**default_sources("excel"), "metrics": []})
    assert not got.errors, got.errors
    assert got.rows["tables"][0]["business_name"] == "Ürün" and got.rows["tables"][0]["row_count"] == 606
    assert got.rows["columns"] == [{**{k: None for k in got.rows["columns"][0]}, "table_name": "dbo.DimProduct",
                                    "column_name": "EnglishProductName", "is_pii": False, "synonyms": "ürün adı, isim"}]
    assert got.rows["relationships"][0]["is_active"] is False
    # sayfa adı büyük/küçük harf duyarsız
    assert not collect(r, {"tables": ["tablolar"], "columns": ["KOLONLAR"], "relationships": [], "metrics": []}).errors
    assert "sayfası bulunamadı" in " ".join(collect(r, {"tables": ["Yok"], "columns": ["Kolonlar"]}).errors)


def test_excel_missing_file():
    with pytest.raises(FileNotFoundError):
        ExcelReader("D:/yok/sozluk.xlsx")


def test_suggest_role_and_defaults():
    assert suggest_role({"table_name", "column_name", "data_type"}) == "columns"
    assert suggest_role({"table_name", "business_name", "subject_area"}) == "tables"
    assert suggest_role({"from_table", "from_column", "to_table", "to_column"}) == "relationships"
    assert suggest_role({"foo"}) is None
    assert DEFAULTS["sqlserver"]["columns"] == ["meta.dd_columns"] and default_sources("mysql")["tables"] == ["dd_tables"]


def test_template_roundtrip(settings, services, tmp_path):
    """Sözlük şablonu indirilip Excel kaynağı olarak geri okunabilmeli (aynı tablo / kolon sayıları)."""
    from fastapi.testclient import TestClient

    import app.main as m

    m.state.services = services
    r = TestClient(m.app).get("/api/settings/dictionary/template.xlsx")
    assert r.status_code == 200 and r.headers["content-type"].startswith("application/vnd.openxmlformats")
    path = tmp_path / "t.xlsx"
    path.write_bytes(r.content)
    got = collect(ExcelReader(path), default_sources("excel"))
    assert not got.errors, got.errors
    dd = services.dictionary
    assert got.counts["tables"] == len([t for t in dd.tables.values() if t.table_type != "view"])
    assert got.counts["columns"] == sum(len(t.columns) for t in dd.tables.values() if t.table_type != "view")
