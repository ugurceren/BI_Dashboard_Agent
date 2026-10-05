"""Sözlükteki tablo / view adlarının veritabanı kataloğuna eşlenmesi (Excel'de farklı yazılmış adlar)."""

import openpyxl

from app.data.connector import QueryResult
from app.dictionary.names import NameResolver, canon, split_name
from app.dictionary.repository import DataDictionary
from app.dictionary.sources import ExcelReader, collect


def test_split_and_canon():
    assert split_name("[EDWDM].[dbo].[vw_Satis]") == ["EDWDM", "dbo", "vw_Satis"]
    assert split_name('"dbo"."Tablo.Adı"') == ["dbo", "Tablo.Adı"]
    assert canon("  EDWDM.dbo.vw_Satis ") == "dbo.vw_Satis"
    assert canon("SUNUCU.EDWDM.rpt.v_Ozet") == "rpt.v_Ozet"
    assert canon("[vw_Satis]") == "vw_Satis"


def test_resolver_maps_to_catalog_objects():
    r = NameResolver([("dbo", "FactSales"), ("rpt", "vw_Satis"), ("dbo", "Ortak"), ("stg", "Ortak")])
    assert r.resolve("dbo.factsales") == "dbo.FactSales"           # yalnız harf farkı: değişiklik sayılmaz
    assert r.resolve("[EDWDM].[rpt].[vw_satis]") == "rpt.vw_Satis"
    assert r.resolve("vw_Satis") == "rpt.vw_Satis"                  # şemasız, katalogda tek
    assert r.resolve("dbo.vw_Satis") == "rpt.vw_Satis"              # yanlış şema, katalogda tek
    assert r.resolve("Ortak") == "dbo.Ortak"                        # iki şemada var: dbo tercih
    assert r.resolve("x.Ortak") == "x.Ortak" and "x.ortak" in r.missing   # belirsiz: eşlenmez
    assert r.resolve("Yok") == "dbo.Yok" and "dbo.yok" in r.missing
    assert "dbo.factsales" not in r.changed and r.changed["vw_Satis"] == "rpt.vw_Satis"


class CatalogConnector:
    """Yalnız katalog sorgularını yanıtlar (sys.objects / sys.columns)."""

    dialect = "tsql"

    def __init__(self, objects: dict[str, tuple[str, list[str]]]):
        self.objects = objects   # "şema.ad" → (tip U/V, kolonlar)

    def execute(self, sql: str, max_rows: int) -> QueryResult:
        if "FROM sys.columns" in sql:
            rows = [[*k.split("."), c, "int", None] for k, (_, cols) in self.objects.items() for c in cols]
            return QueryResult(["s", "o", "c", "t", "d"], ["string"] * 5, rows, False, 1)
        if "FROM sys.objects" in sql:
            rows = [[*k.split("."), t, 1, None, 10] for k, (t, _) in self.objects.items()]
            return QueryResult(["s", "o", "t", "p", "d", "r"], ["string"] * 6, rows, False, 1)
        return QueryResult([], [], [], False, 1)


def test_excel_names_resolved_against_catalog(settings, tmp_path, monkeypatch):
    """Excel'de şema ayrı kolonda / şemasız / veritabanı önekli yazılan view'lar sözlükte TANIMLI görünür."""
    p = tmp_path / "sozluk.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Kolonlar"
    ws.append(["TABLE_SCHEMA", "TABLE_NAME", "COLUMN_NAME", "Açıklama"])
    ws.append(["rpt", "vw_Satis", "Tutar", "Satış tutarı"])
    ws.append([None, "vw_Musteri", "MusteriKey", "Müşteri"])            # şemasız
    ws.append([None, "[EDWDM].[dbo].[FactSales]", "Amount", "Tutar"])   # veritabanı önekli, parantezli
    ws.append([None, "dbo.Silinmis", "X", ""])                          # veritabanında yok
    wb.save(p)

    got = collect(ExcelReader(p), {"tables": [], "columns": ["Kolonlar"], "relationships": [], "metrics": []})
    assert not got.errors and {r["table_name"] for r in got.rows["columns"]} >= {"rpt.vw_Satis"}

    con = CatalogConnector({"rpt.vw_Satis": ("V", ["Tutar", "Bolge"]), "crm.vw_Musteri": ("V", ["MusteriKey"]),
                            "dbo.FactSales": ("U", ["Amount"])})
    import app.data.connections as conns
    import app.dictionary.sources as dsrc
    monkeypatch.setattr(conns, "saved_dictionary", lambda: ({"kind": "excel", "excel_path": str(p)}, {"columns": ["Kolonlar"]}))
    monkeypatch.setattr(conns, "open_dictionary_reader", lambda f: (ExcelReader(p), {}))
    monkeypatch.setattr(dsrc, "collect", collect)
    monkeypatch.setattr(DataDictionary, "_view_registry", lambda self: [])
    dd = DataDictionary(settings, con).load()

    for name in ("rpt.vw_satis", "crm.vw_musteri", "dbo.factsales"):
        t = dd.tables[name]
        assert t.documented and t.in_db and t.can_select, name
    assert dd.tables["rpt.vw_satis"].display_name == "rpt.vw_Satis"
    assert {c.name for c in dd.tables["rpt.vw_satis"].columns} == {"tutar", "bolge"}   # eksik kolon katalogdan
    assert not [t for t in dd.tables.values() if not t.documented]                     # kopya "tanımsız" yok
    assert dd.missing_in_db == ["dbo.Silinmis"]
    assert dd.name_changes["vw_Musteri"] == "crm.vw_Musteri"
    assert dd.is_view("rpt.vw_satis") and not dd.is_view("dbo.factsales") and dd.tables["dbo.factsales"].object_type == "table"


def test_corporate_excel_layout_view_name_and_camelcase_headers(settings, tmp_path, monkeypatch):
    """Kurum sözlüğü: 'column' ve 'table' sayfaları; DatabaseName / SchemaName / view_name / ColumnName / ColumnDescription."""
    from app.dictionary.sources import _norm_key, suggest_role

    assert [_norm_key(h) for h in ["DatabaseName", "SchemaName", "view_name", "ColumnName", "ColumnDescription", "Description"]] == \
        ["database_name", "schema_name", "table_name", "column_name", "description", "description"]
    assert _norm_key("IsPII") == "is_pii" and _norm_key("Tablo Adı") == "table_name" and _norm_key("KolonAdı") == "column_name"

    p = tmp_path / "veri sözlüğü.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "column"
    ws.append(["DatabaseName", "SchemaName", "view_name", "ColumnName", "ColumnDescription"])
    ws.append(["EDWDM", "CLT", "vRepurchaseGuarantee", "CollateralName", "Teminat Adı"])
    ws.append(["EDWDM", "CLT", "vRepurchaseGuarantee", "CustomerPartyId", "Müşteri Numarası (SKEY)"])
    ws.append(["EDWDM", "CLT", "vRepurchaseGuaranteeMasked", "CustomerName", "Müşteri Adı"])
    t = wb.create_sheet("table")
    t.append(["DatabaseName", "SchemaName", "view_name", "Description"])
    t.append(["EDWDM", "CLT", "vRepurchaseGuarantee", "Geri alım garantisi ile ilgili bilgileri içermektedir."])
    wb.save(p)

    reader = ExcelReader(p)
    roles = {name: suggest_role({_norm_key(c) for c in cols}) for name, cols in reader.list_tables().items()}
    assert roles == {"column": "columns", "table": "tables"}    # Excel yüklenince sayfalar otomatik seçilir
    got = collect(reader, {"tables": ["table"], "columns": ["column"], "relationships": [], "metrics": []})
    assert not got.errors, got.errors
    assert {r["table_name"] for r in got.rows["columns"]} == {"CLT.vRepurchaseGuarantee", "CLT.vRepurchaseGuaranteeMasked"}

    con = CatalogConnector({"CLT.vRepurchaseGuarantee": ("V", ["CollateralName", "CustomerPartyId"]),
                            "CLT.vRepurchaseGuaranteeMasked": ("V", ["CustomerName"]),
                            "CLT.vRepurchaseGuaranteePersonnelMasked": ("V", ["X"])})
    import app.data.connections as conns
    monkeypatch.setattr(conns, "saved_dictionary", lambda: ({"kind": "excel"}, {"tables": ["table"], "columns": ["column"]}))
    monkeypatch.setattr(conns, "open_dictionary_reader", lambda f: (ExcelReader(p), {}))
    monkeypatch.setattr(DataDictionary, "_view_registry", lambda self: [])
    dd = DataDictionary(settings, con).load()
    v = dd.tables["clt.vrepurchaseguarantee"]
    assert v.documented and v.in_db and v.description.startswith("Geri alım garantisi")
    assert {c.name: c.description for c in v.columns}["collateralname"] == "Teminat Adı"
    assert dd.tables["clt.vrepurchaseguaranteemasked"].documented
    assert not dd.tables["clt.vrepurchaseguaranteepersonnelmasked"].documented   # Excel'de yok: "sözlükte yok" doğru
    # sözlükte tablo / view ayrımı yok: tip veritabanı kataloğundan (sys.objects.type = 'V')
    assert v.object_type == "view" and dd.is_view("clt.vrepurchaseguarantee")
    assert dd.is_view("clt.vrepurchaseguaranteepersonnelmasked")


def _abc_excel(tmp_path):
    p = tmp_path / "abc.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sayfa1"
    ws.append(["abc", "sema", "xyz", "notlar"])      # tablo adı / kolon adı başlıkları tanınmaz
    ws.append(["vRepurchaseGuarantee", "CLT", "CollateralName", "Teminat Adı"])
    ws.append(["vRepurchaseGuarantee", "CLT", "CustomerPartyId", "Müşteri No"])
    ws.append(["vRepurchaseGuaranteeMasked", "CLT", "CustomerName", "Müşteri Adı"])
    wb.save(p)
    return p


def test_unrecognized_headers_detected_from_content(tmp_path):
    """Başlık 'abc' / 'xyz' olsa da değerleri veritabanındaki view / kolon adlarıyla örtüştüğü için bulunur."""
    from app.dictionary.sources import Known

    known = Known.from_catalog([("CLT", "vRepurchaseGuarantee"), ("CLT", "vRepurchaseGuaranteeMasked")],
                               ["CollateralName", "CustomerPartyId", "CustomerName"])
    src = {"tables": [], "columns": ["Sayfa1"], "relationships": [], "metrics": []}
    plain = collect(ExcelReader(_abc_excel(tmp_path)), src)                # katalog yok: tanınamaz, açık hata
    assert plain.errors and "Başlık eşleme" in plain.errors[0]
    assert plain.sources["Sayfa1"]["headers"] == ["abc", "sema", "xyz", "notlar"]

    got = collect(ExcelReader(_abc_excel(tmp_path)), src, known=known)
    assert not got.errors, got.errors
    info = got.sources["Sayfa1"]
    assert info["mapping"]["table_name"] == "abc" and info["how"]["table_name"] == "content"
    assert info["mapping"]["column_name"] == "xyz" and info["how"]["column_name"] == "content"
    assert info["mapping"]["schema_name"] == "sema"                       # 'sema' başlıktan (Türkçe karşılık)
    assert {r["table_name"] for r in got.rows["columns"]} == {"CLT.vRepurchaseGuarantee", "CLT.vRepurchaseGuaranteeMasked"}
    assert got.learned_mappings() == {"Sayfa1": {"table_name": "abc", "column_name": "xyz"}}


def test_manual_mapping_overrides_and_disables(tmp_path):
    src = {"tables": [], "columns": ["Sayfa1"], "relationships": [], "metrics": []}
    manual = {"sayfa1": {"table_name": "abc", "column_name": "xyz", "description": "notlar", "schema_name": ""}}
    got = collect(ExcelReader(_abc_excel(tmp_path)), src, manual)        # katalog olmadan da elle eşlemeyle çalışır
    assert not got.errors, got.errors
    info = got.sources["Sayfa1"]
    assert info["how"]["table_name"] == "manual" and info["mapping"]["schema_name"] is None   # "" = kullanma
    row = got.rows["columns"][0]
    assert row["table_name"] == "vRepurchaseGuarantee" and row["description"] == "Teminat Adı"


def test_settings_test_returns_mapping_and_save_remembers_detection(settings, services, conn_file, monkeypatch, tmp_path):
    from fastapi.testclient import TestClient

    import app.data.connections as conns
    import app.main as m
    from app.dictionary.sources import Known

    m.state.services = services
    m.state.agent = type("A", (), {"services": services})()
    monkeypatch.setattr(m, "build_services", lambda: services)
    known = Known.from_catalog([("CLT", "vRepurchaseGuarantee"), ("CLT", "vRepurchaseGuaranteeMasked")],
                               ["CollateralName", "CustomerPartyId", "CustomerName"])
    monkeypatch.setattr(m, "_data_catalog", lambda data: {"known": known})
    c = TestClient(m.app)
    p = _abc_excel(tmp_path)
    dic = {"kind": "excel", "excel_path": str(p), "sources": {"tables": [], "columns": ["Sayfa1"], "relationships": [], "metrics": []}}
    data = {"server": "PASIFIK", "database": "EDWDM"}
    t = c.post("/api/settings/connections/test", json={"target": "dictionary", "data": data, "dictionary": dic}).json()
    assert t["ok"], t
    info = t["sources_info"]["Sayfa1"]
    assert info["how"] == {"table_name": "content", "column_name": "content", "schema_name": "header"}
    # elle: açıklama 'notlar' sütunundan
    dic["mappings"] = {"Sayfa1": {"description": "notlar"}}
    assert c.put("/api/settings/connections/dictionary", json={"dictionary": dic, "data": data}).status_code == 200
    saved = conns.load_connections()["dictionary"]["mappings"]["Sayfa1"]
    assert saved == {"description": "notlar", "table_name": "abc", "column_name": "xyz"}
    # eşleme kayıtlı: katalog olmadan da yüklenir
    monkeypatch.setattr(m, "_data_catalog", lambda data: {})
    assert conns.saved_dictionary()[0]["mappings"]["Sayfa1"]["table_name"] == "abc"



def test_agent_sees_only_usable_tables_and_views(services):
    """Sözlükte olup veritabanında olmayan (ör. sözlüğe yazılmış SP) ya da SELECT yetkisi olmayan nesneler
    agent'ın search_dictionary / get_table_details araçlarında görünmez."""
    from app.dictionary.repository import DDColumn, DDTable
    from app.harness.session import Session
    from app.harness.tools import ToolContext, h_get_table_details, h_search_dictionary

    dd = services.dictionary
    sp = DDTable("clt.usprepurchaseguarantee", "Geri Alım Garantisi Prosedürü", "Geri alım garantisi SP", "Kredi", "", None,
                 display_name="CLT.uspRepurchaseGuarantee", in_db=False)
    sp.columns = [DDColumn(sp.name, "guaranteeamount", "Garanti Tutarı", "", "", "measure", "sum", [], False, "")]
    locked = DDTable("clt.vrepurchaseguaranteemasked", "Geri Alım Garantisi (maskeli)", "Geri alım garantisi", "Kredi", "", 5,
                     display_name="CLT.vRepurchaseGuaranteeMasked", in_db=True, can_select=False)
    locked.columns = [DDColumn(locked.name, "guaranteeamount", "Garanti Tutarı", "", "", "measure", "sum", [], False, "")]
    ok = DDTable("clt.vrepurchaseguarantee", "Geri Alım Garantisi", "Geri alım garantisi", "Kredi", "", 5,
                 display_name="CLT.vRepurchaseGuarantee", in_db=True, can_select=True)
    ok.columns = [DDColumn(ok.name, "guaranteeamount", "Garanti Tutarı", "", "", "measure", "sum", [], False, "")]
    for x in (sp, locked, ok):
        dd.tables[x.name] = x
    try:
        ctx = ToolContext(Session(), services)
        found = [h["table"] for h in h_search_dictionary(ctx, {"query": "geri alım garantisi"}).content["tables"]]
        assert "clt.vrepurchaseguarantee" in found
        assert "clt.usprepurchaseguarantee" not in found and "clt.vrepurchaseguaranteemasked" not in found
        r = h_get_table_details(ctx, {"tables": ["CLT.uspRepurchaseGuarantee", "CLT.vRepurchaseGuarantee"]})
        assert [d["table"] for d in r.content["tables"]] == ["clt.vrepurchaseguarantee"]
        assert r.content["not_found"] == ["CLT.uspRepurchaseGuarantee"]
    finally:
        for x in (sp, locked, ok):
            dd.tables.pop(x.name, None)
