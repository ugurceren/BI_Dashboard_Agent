"""Birden çok veritabanı (aynı sunucu): birincil (bağlı) veritabanı şema.nesne, ek veritabanları db.şema.nesne.
Kullanıcı yalnız SELECT yetkisi olan nesneleri görür; seçili olmayan veritabanları reddedilir; veritabanına özgü kurallar
(NOLOCK, DataDate, yetki varyantları, kural dosyaları) nesnenin kendi veritabanına göre uygulanır."""

import pytest

from app.data.connector import QueryResult
from app.data.validator import RolePolicy, SqlValidator
from app.dictionary.names import NameResolver, canon
from app.dictionary.repository import DataDictionary
from app.harness.rules import databases_block, phase_rules
from app.harness.session import Session

POL = RolePolicy("standart", ["*"])


class Cat:
    """Tek bir veritabanının katalog bağlantısı (DB_NAME + sys.objects / sys.columns)."""

    dialect = "tsql"

    def __init__(self, db, objects, denied=()):
        self.db, self.objects, self.denied, self.executed = db, objects, set(denied), []

    def execute(self, sql, max_rows):
        self.executed.append(sql)
        if sql.startswith("SELECT DB_NAME()"):
            return QueryResult(["db"], ["string"], [[self.db]], False, 1)
        if "FROM sys.columns" in sql:
            rows = [[*k.split("."), c, t, None] for k, cols in self.objects.items() for c, t in cols]
            return QueryResult(list("sotcd"), ["string"] * 5, rows, False, 1)
        if "FROM sys.objects" in sql:
            rows = [[*k.split("."), "V", 0 if k in self.denied else 1, None, 10] for k in self.objects]
            return QueryResult(list("sotpdr"), ["string"] * 6, rows, False, 1)
        return QueryResult([], [], [], False, 1)


@pytest.fixture()
def multi(settings, monkeypatch):
    import app.data.connections as conns
    monkeypatch.setattr(conns, "saved_dictionary", lambda: None)
    monkeypatch.setattr(DataDictionary, "_view_registry", lambda self: [])
    monkeypatch.setattr(DataDictionary, "_connector", lambda self, cfg: Cat("x", {}))   # dictionary.toml sözlüğü boş
    primary = Cat("EDWDM", {"CLT.vGuarantee": [("DataDate", "date"), ("CustomerPartyId", "int"), ("Amount", "decimal")]})
    edw = Cat("EDW", {"dbo.FactLoan": [("DataDate", "date"), ("CustomerPartyId", "int"), ("Balance", "decimal")],
                      "dbo.Secret": [("X", "int")]}, denied={"dbo.Secret"})
    dd = DataDictionary(settings, primary)
    dd.extra_databases = ["EDW"]
    dd.extra_connector = lambda db: {"EDW": edw}[db]
    monkeypatch.setattr(DataDictionary, "load", DataDictionary.load)
    # dictionary.toml'daki sorgular yerine boş sözlük: yalnız katalog
    monkeypatch.setattr("app.dictionary.repository.load_toml", lambda p: {"queries": {}, "snapshot_date_columns": ["DataDate"],
                                                                         "snapshot_databases": ["EDWDM"],
                                                                         "view_variant_databases": ["EDWDM"],
                                                                         "nolock_databases": ["EDWDM", "EDW"]})
    dd.load()
    return dd, primary, edw


def test_catalog_reads_each_database_with_own_permissions(multi):
    dd, primary, edw = multi
    assert dd.database == "EDWDM" and dd.databases() == ["EDWDM", "EDW"]
    assert "clt.vguarantee" in dd.tables and "edw.dbo.factloan" in dd.tables
    assert dd.tables["edw.dbo.factloan"].display_name == "EDW.dbo.FactLoan"
    assert "edw.dbo.secret" not in dd.tables                                  # SELECT yetkisi yok: listelenmez
    assert any("HAS_PERMS_BY_NAME" in q for q in edw.executed)               # yetki EDW'nin kendi bağlamında
    assert dd.db_of("edw.dbo.factloan") == "EDW" and dd.db_of("clt.vguarantee") == "EDWDM"


def test_rules_follow_each_objects_own_database(multi):
    dd, *_ = multi
    assert dd.tables["clt.vguarantee"].snapshot_date == "DataDate"          # EDWDM: günlük anlık görüntü
    assert not dd.tables["edw.dbo.factloan"].snapshot_date                 # EDW: kural yok (aynı adlı kolon olsa da)
    assert dd.nolock("clt.vguarantee") and dd.nolock("edw.dbo.factloan")   # veri ambarı: ikisinde de NOLOCK


@pytest.mark.parametrize("sql,ok", [
    ("SELECT SUM(l.Balance) FROM EDW.dbo.FactLoan l WHERE l.CustomerPartyId > 0", True),
    ("SELECT SUM(l.Balance) FROM edw.dbo.factloan l WHERE l.CustomerPartyId > 0", True),       # harf duyarsız
    ("SELECT g.CustomerPartyId, SUM(l.Balance) FROM CLT.vGuarantee g JOIN EDW.dbo.FactLoan l "
     "ON l.CustomerPartyId = g.CustomerPartyId WHERE g.DataDate = (SELECT MAX(DataDate) FROM CLT.vGuarantee WHERE DataDate >= DATEADD(day, -2, CAST(GETDATE() AS DATE))) "
     "GROUP BY g.CustomerPartyId", True),                                                      # veritabanları arası join
    ("SELECT TOP 5 CustomerPartyId FROM EDWDM.CLT.vGuarantee WHERE DataDate = "
     "(SELECT MAX(DataDate) FROM EDWDM.CLT.vGuarantee WHERE DataDate >= DATEADD(day, -2, CAST(GETDATE() AS DATE)))", True),                                 # birincil öneki atılır
    ("SELECT TOP 5 CustomerPartyId FROM EDWDM.CLT.vGuarantee WHERE CustomerPartyId > 0", False),  # anlık görüntü: gün yok
    ("SELECT TOP 5 X FROM EDW.dbo.Secret WHERE X > 0", False),                                 # yetkisiz nesne
    ("SELECT TOP 5 name FROM master.sys.objects WHERE object_id > 0", False),                  # seçili olmayan veritabanı
    ("SELECT TOP 5 Balance FROM SRV.EDW.dbo.FactLoan WHERE Balance > 0", False),               # sunucu adı
])
def test_validator_with_multiple_databases(multi, sql, ok):
    dd, *_ = multi
    r = SqlValidator(dd, "tsql", []).validate(sql, POL)
    assert r.ok is ok, r.errors
    if ok:
        assert "EDWDM." not in r.sql                                          # birincil veritabanı adı yazılmaz
        assert r.sql.upper().count("WITH (NOLOCK)") == r.sql.upper().count(" FROM ") + r.sql.upper().count(" JOIN ")


def test_policy_schema_check_uses_schema_part(multi):
    dd, *_ = multi
    pol = RolePolicy("standart", ["*"], denied_schemas=["dbo"])
    r = SqlValidator(dd, "tsql", []).validate("SELECT TOP 5 Balance FROM EDW.dbo.FactLoan WHERE Balance > 0", pol)
    assert not r.ok and "dbo" in r.errors[0]                               # şema = dbo (veritabanı adı değil)
    assert pol.denial_reason("edw.dbo.factloan") and not pol.denial_reason("clt.vguarantee")


def test_name_resolution_keeps_selected_extra_database():
    assert canon("EDW.dbo.FactLoan", ["EDW"]) == "EDW.dbo.FactLoan"
    assert canon("EDWDM.CLT.vGuarantee", ["EDW"]) == "CLT.vGuarantee"       # birincil: önek atılır
    r = NameResolver([("CLT", "vGuarantee"), ("EDW.dbo", "FactLoan"), ("dbo", "Ortak"), ("EDW.dbo", "Ortak")], ["EDW"])
    assert r.resolve("edw.dbo.factloan") == "EDW.dbo.FactLoan"
    assert r.resolve("FactLoan") == "EDW.dbo.FactLoan"                       # veritabanı yazılmamış, tek eşleşme
    assert r.resolve("Ortak") == "dbo.Ortak"                                 # iki veritabanında: önce birincil
    assert r.resolve("EDW.dbo.Ortak") == "EDW.dbo.Ortak"
    assert r.resolve("EDW.dbo.vGuarantee") == "EDW.dbo.vGuarantee" and "edw.dbo.vguarantee" in r.missing  # başka db'ye atlamaz


def test_excel_database_name_column(tmp_path):
    import openpyxl

    from app.dictionary.sources import ExcelReader, collect
    p = tmp_path / "s.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "column"
    ws.append(["DatabaseName", "SchemaName", "view_name", "ColumnName"])
    ws.append(["EDW", "dbo", "FactLoan", "Balance"])
    ws.append(["EDWDM", "CLT", "vGuarantee", "Amount"])
    wb.save(p)
    got = collect(ExcelReader(p), {"tables": [], "columns": ["column"], "relationships": [], "metrics": []})
    assert {r["table_name"] for r in got.rows["columns"]} == {"EDW.dbo.FactLoan", "EDWDM.CLT.vGuarantee"}


def test_llm_is_told_about_databases(multi, services):
    dd, *_ = multi
    block = databases_block(dd)
    assert "EDWDM (bağlı / birincil)" in block and "EDW.şema.nesne" in block
    from dataclasses import replace
    svc = replace(services, dictionary=dd)
    data = phase_rules(svc, Session(phase="data"))
    assert "Veri kaynağı veritabanları" in data and "SQL kullanım standartları (geçerli veritabanları: EDWDM, EDW)" in data
    assert "CustomerPartyId" in data and "AccountNumber" in data
    assert "Kurum veri kuralları — EDWDM (yalnız EDWDM nesneleri için" in data


def test_extra_database_failure_does_not_break_primary(settings, monkeypatch):
    import app.data.connections as conns
    monkeypatch.setattr(conns, "saved_dictionary", lambda: None)
    monkeypatch.setattr(DataDictionary, "_view_registry", lambda self: [])
    monkeypatch.setattr(DataDictionary, "_connector", lambda self, cfg: Cat("x", {}))
    monkeypatch.setattr("app.dictionary.repository.load_toml", lambda p: {"queries": {}})
    dd = DataDictionary(settings, Cat("EDWDM", {"CLT.vGuarantee": [("Amount", "decimal")]}))
    dd.extra_databases = ["EDW"]

    def boom(db):
        raise RuntimeError("Cannot open database \"EDW\" requested by the login.")
    dd.extra_connector = boom
    dd.load()
    assert "clt.vguarantee" in dd.tables and "EDW" in dd.extra_errors
    assert "OKUNAMIYOR" in databases_block(dd)
