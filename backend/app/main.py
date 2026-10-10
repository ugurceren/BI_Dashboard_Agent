"""FastAPI uygulaması. Çalıştırma:  uvicorn app.main:app --port 8000 --reload"""

from __future__ import annotations

import hashlib
import json
import logging
import os
import re
import threading
import time
from contextlib import asynccontextmanager
from pathlib import Path
from typing import Any, Iterator
from urllib.parse import quote

# Kurum içi HTTPS (LLM geçidi vb.): Python kendi sertifika listesini değil Windows sertifika deposunu kullansın;
# böylece kurumun kendi sertifika otoritesiyle imzalanmış adresler tarayıcıdaki gibi tanınır. Kapatmak: BI_SYSTEM_CERTS=0
if os.environ.get("BI_SYSTEM_CERTS", "1") != "0":
    try:
        import truststore

        truststore.inject_into_ssl()
    except Exception:  # noqa: BLE001 — truststore yoksa Python'un kendi listesiyle devam
        pass

from fastapi import Body, FastAPI, HTTPException, Request
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import HTMLResponse, JSONResponse, Response, StreamingResponse
from pydantic import BaseModel, ValidationError, field_validator
from sqlglot import exp

from app.config import BACKEND_DIR, get_settings, load_toml
from app.data.connector import QueryError, create_connector
from app.data.model_filters import ModelFilter, ModelFilterEngine, with_nolock
from app.data.views import ViewRegistry, build_view_script, safe_view_name, select_from_view, view_columns
from app.data.validator import RolePolicy, SqlValidator
from app.dictionary.repository import DataDictionary
from app import authz
from app.identity import Identity, current_identity
from app.meta.store import MetaStore
from app.harness.agent import Agent, Audit, Event
from app.harness.session import PHASES, QueryDraft, Requirements, SessionStore, TranscriptItem, now_iso
from app.harness.tools import (_DESIGN_KICKOFF, Services, ToolContext, _build_dataset, _run_validated, _short_db_error,
                               _validate_spec, add_join_relationships)
from app.llm.gateway import LLMGateway


def conns_llm_settings(settings):
    from app.data.connections import llm_settings
    return llm_settings(settings)
from app.spec.models import Dataset, DatasetField, ReportSpec

logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(name)s: %(message)s")
log = logging.getLogger("bi-agent")



class AppState:
    services: Services
    gateway: LLMGateway
    store: SessionStore
    agent: Agent
    startup_error: str | None = None
    meta: MetaStore | None = None          # platform kayıtları (Vitrin)
    meta_error: str | None = None


class _BrokenConnector:
    """Bağlantı cümlesi kurulamadığında yer tutucu: her sorguda anlaşılır hata verir."""
    dialect = "tsql"

    def __init__(self, error: str):
        self.error = error

    def execute(self, sql: str, max_rows: int):
        raise QueryError(self.error)

    def ping(self) -> None:
        raise QueryError(self.error)


state = AppState()


def build_services() -> Services:
    """Servisler. Veritabanı / sözlük erişilemezse uygulama yine açılır (boş sözlükle); hata state.startup_error'da
    tutulur ve arayüzde gösterilir — kullanıcı Bağlantı Ayarları'ndan düzeltebilir."""
    settings = get_settings()
    state.startup_error = None
    try:
        connector = create_connector(settings)
    except Exception as e:  # noqa: BLE001
        log.error("Veri bağlantısı kurulamadı: %s", e)
        state.startup_error = f"Veri bağlantısı: {e}"
        connector = _BrokenConnector(str(e))
    dictionary = DataDictionary(settings, connector)
    dictionary.extra_databases = conns.extra_databases(settings)   # aynı sunucudaki ek veritabanları (ör. EDW)
    try:
        dictionary.load()
    except Exception as e:  # noqa: BLE001
        from app.data.odbc import friendly_error
        msg = friendly_error(str(e))
        log.error("Veri sözlüğü yüklenemedi: %s", msg)
        state.startup_error = (state.startup_error + " | " if state.startup_error else "") + f"Veri sözlüğü: {msg}"
        # sözlük olmasa da veritabanındaki yetkili tablo / view'lar görünsün ve sorgulanabilsin
        try:
            dictionary._merge_catalog()
            dictionary.mark_snapshots()
            dictionary.apply_model_relationships()
        except Exception as ce:  # noqa: BLE001
            log.warning("Veritabanı kataloğu da okunamadı: %s", ce)
    policy_cfg = load_toml(settings.policy_config)
    policies = {name: RolePolicy(name=name, **cfg) for name, cfg in policy_cfg.get("roles", {}).items()}
    validator = SqlValidator(dictionary, connector.dialect, policy_cfg.get("sql", {}).get("denied_functions", []))
    return Services(settings, dictionary, connector, validator, policies)


def build_meta() -> MetaStore | None:
    """Platform meta deposu: SQL Server (META_ODBC ya da Bağlantı Ayarları → Platform Veritabanı), yoksa yerel SQLite."""
    settings = get_settings()
    state.meta_error = None
    odbc = conns.meta_odbc(settings)
    if odbc:
        try:
            st = MetaStore.sqlserver(odbc)
            log.info("Platform veritabanı: SQL Server")
            return st
        except Exception as e:  # noqa: BLE001
            from app.data.odbc import friendly_error
            state.meta_error = f"Platform veritabanına bağlanılamadı: {friendly_error(str(e))[:300]}"
            log.error(state.meta_error)
            if settings.platform_mode == "server":
                return None    # sunucuda sessizce yerel dosyaya düşme: kayıtlar bölünmesin
    try:
        return MetaStore.sqlite(settings.meta_sqlite)
    except Exception as e:  # noqa: BLE001
        state.meta_error = f"Platform kayıt dosyası açılamadı: {e}"
        log.error(state.meta_error)
        return None


@asynccontextmanager
async def lifespan(_: FastAPI):
    settings = get_settings()
    state.services = build_services()
    state.gateway = LLMGateway(conns_llm_settings(settings))
    state.store = SessionStore(settings.sessions_dir)
    removed = state.store.dedupe_titles()
    if removed:
        log.info("Aynı başlıklı %d eski rapor oturumu kaldırıldı (aynı isimle tek rapor).", removed)
    state.agent = Agent(state.gateway, state.services, state.store)
    state.meta = build_meta()
    log.info("Sözlük: %d tablo, %d metrik | LLM: %s @ %s | vision: %s", len(state.services.dictionary.tables),
             len(state.services.dictionary.metrics), state.gateway.s.llm_model, state.gateway.s.llm_base_url,
             state.gateway.s.vision_model or "-")
    yield


app = FastAPI(title="BI Lens", lifespan=lifespan)
app.add_middleware(CORSMiddleware, allow_origins=get_settings().cors_origins, allow_methods=["*"], allow_headers=["*"])


def _session(sid: str):
    try:
        return state.store.get(sid)
    except KeyError:
        raise HTTPException(404, "Oturum bulunamadı")


# --------------------------------------------------------------------------- sağlık
@app.get("/api/health")
def health() -> dict[str, Any]:
    s = get_settings()
    data: dict[str, Any] = {"ok": True, "dialect": state.services.connector.dialect}
    try:
        state.services.connector.ping()
    except Exception as e:  # noqa: BLE001
        data = {"ok": False, "dialect": state.services.connector.dialect, "error": str(e)[:200]}
    llm = state.gateway.health()
    dic = {"tables": len(state.services.dictionary.tables)}
    if state.startup_error and "sözlüğü" in state.startup_error:
        dic["error"] = state.startup_error.split("Veri sözlüğü: ", 1)[-1][:300]
    if data["ok"] is False:
        from app.data.odbc import friendly_error
        data["error"] = friendly_error(data.get("error", ""))[:400]
    return {"ok": llm["reachable"] and data["ok"], "llm": llm,
            "vision": {"configured": bool(state.gateway.s.vision_model), "model": state.gateway.s.vision_model}, "data": data,
            "dictionary": dic}


# --------------------------------------------------------------------------- oturumlar
@app.get("/api/sessions")
def list_sessions(request: Request) -> list[dict[str, Any]]:
    items = _visible(_need(request, "builder"), state.store.list())
    published = {r["session_id"]: r for r in state.meta.reports()} if state.meta is not None else {}
    kinds = state.services.dictionary._table_kinds()
    for it in items:
        it["domains"] = _report_domains(it.get("source_tables") or [], kinds)
        rep = published.get(it["id"])
        if rep:
            it["published"] = {"report_id": rep["report_id"], "version": rep["current_version"], "status": rep["status"],
                               "updated_at": rep["updated_at"]}
    return items


def _report_domains(tables: list[str], kinds: dict[str, str]) -> list[str]:
    """Raporun iş alanı (domain): kullandığı tabloların sözlükteki konu alanı.
    Olgu (fact) tablolarının alanı önce gelir; onaylı view'lar kaynak tablolarına çözülür;
    'Ortak' (tarih, para birimi…) yalnızca başka alan yoksa gösterilir."""
    repo = state.services.dictionary
    order = {"fact": 0, "bridge": 1, "dimension": 2}
    ranked: list[tuple[int, str]] = []
    for name in tables:
        t = repo.tables.get(name)
        if t is None:
            continue
        if kinds.get(name) == "view":  # kolon soy ağacındaki kaynak tabloların alanı
            src = {v.rsplit(".", 1)[0] for k, v in repo.view_lineage.items() if k.startswith(name + ".")}
            ranked += [(order.get(kinds.get(x, ""), 2), repo.tables[x].subject_area) for x in src if x in repo.tables]
        elif t.subject_area:
            ranked.append((order.get(kinds.get(name, ""), 2), t.subject_area.strip()))
    out: list[str] = []
    for _, a in sorted(ranked, key=lambda x: x[0]):
        if a and a not in out:
            out.append(a)
    main = [a for a in out if a.lower() != "ortak"]
    return main or out


def _me(request: Request) -> Identity:
    """Bağlanan kullanıcı + platform rolü. Sunucu modunda kimlik yalnız ters proxy başlığından gelir."""
    settings = get_settings()
    if settings.platform_mode == "server" and not (request.headers.get("x-remote-user") or "").strip():
        raise HTTPException(401, "Kimlik doğrulanamadı: uygulamaya kurumsal oturum açma (IIS / Windows kimlik doğrulaması) "
                                 "üzerinden bağlanın.")
    return authz.resolve(settings, state.meta, current_identity(settings, dict(request.headers)))


def _need(request: Request, role: str) -> Identity:
    """En az bu platform rolü (viewer < builder < admin)."""
    ident = _me(request)
    if not authz.at_least(ident, role):
        names = {"builder": "rapor tasarımcısı (builder)", "admin": "yönetici (admin)"}
        raise HTTPException(403, f"Bu işlem için {names.get(role, role)} yetkisi gerekiyor.")
    return ident


def _own_session(request: Request, sid: str):
    """Tasarım oturumu: yalnız sahibi ya da admin (masaüstü modunda tek kullanıcı → hepsi)."""
    ident = _need(request, "builder")
    s = _session(sid)
    if not authz.at_least(ident, "admin") and not authz.owns(ident, s.owner):
        raise HTTPException(403, "Bu rapor başka bir kullanıcıya ait.")
    return s, ident


def _visible(ident: Identity, items: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """Envanter: admin hepsini, builder yalnız kendi raporlarını görür."""
    if authz.at_least(ident, "admin"):
        return items
    return [it for it in items if authz.owns(ident, it.get("owner"))]


def _meta() -> MetaStore:
    if state.meta is None:
        raise HTTPException(503, state.meta_error or "Platform veritabanı yapılandırılmamış.")
    return state.meta


def _audit(ident: Identity, event: str, report_id: str | None = None, **details: Any) -> None:
    if state.meta is not None:
        state.meta.audit(ident.username, event, report_id, **details)


@app.post("/api/sessions")
def create_session(request: Request) -> dict[str, Any]:
    me = _need(request, "builder")
    return state.store.create(me.role, owner=me.username, owner_name=me.display_name).public()


@app.get("/api/me")
def me(request: Request) -> dict[str, Any]:
    """Bağlanan kullanıcı (Windows oturumu / LDAP) ve rol politikası."""
    ident = _me(request)
    pol = state.services.policy(ident.role)
    caps = {"design": authz.at_least(ident, "builder"), "admin": authz.at_least(ident, "admin"),
            "vitrin": state.meta is not None}
    return {**ident.public(), "capabilities": caps, "policy": {"allowed_schemas": pol.allowed_schemas, "denied_tables": pol.denied_tables,
                                          "allow_pii": pol.allow_pii, "max_rows": pol.max_rows}}


def _object_access(pol: RolePolicy, t) -> tuple[bool, str | None]:
    """Nesne erişimi: önce rol politikası (yasaklı şema / tablo), sonra veritabanının kendi yetkisi (SELECT)."""
    reason = pol.denial_reason(t.name)
    if reason:
        return False, reason
    if t.can_select is False:
        return False, "veritabanında SELECT yetkiniz yok"
    if t.in_db is False:
        return False, "sözlükte var, veritabanında bulunamadı"
    return True, None


@app.get("/api/me/access")
def my_access(request: Request) -> dict[str, Any]:
    """Kullanıcının yetkili olduğu nesneler — tek liste: tablo / view / dataset;
    her biri domain (konu alanı) ve erişim durumuyla (Veri Erişimim sayfası: domain ya da nesne tipine göre gruplanır)."""
    from app.harness.session import _source_tables
    ident = _need(request, "builder")
    pol = state.services.policy(ident.role)
    dd = state.services.dictionary
    kinds = dd._table_kinds()
    registry = {str(v.get("name", "")).lower(): v for v in ViewRegistry(get_settings().views_registry).all()}
    objects: list[dict[str, Any]] = []
    for t in dd.tables.values():
        ok, reason = _object_access(pol, t)
        pii = [c.display_name or c.name for c in t.columns if c.is_pii]
        is_view = dd.is_view(t.name, kinds)   # tablo / view ayrımı veritabanı kataloğundan
        area = t.subject_area
        extra: dict[str, Any] = {}
        if is_view:
            reg = registry.get(t.name, {})
            src = _source_tables([{"sql": reg.get("original_sql")}]) if reg.get("original_sql") else []
            area = ((_report_domains(src, kinds) if src else []) or _report_domains([t.name], kinds) or [area])[0]
            extra = {"report_id": reg.get("session_id"), "dataset_id": reg.get("dataset_id"), "created_at": reg.get("created_at"),
                     "tables": src}
        objects.append({
            "type": "view" if is_view else "table", "id": t.name, "name": t.display_name or t.name,
            "business_name": t.business_name, "description": t.description, "subject_area": area or "Diğer",
            "row_count": t.row_count, "column_count": len(t.columns), "pii_columns": pii,
            "pii_blocked": bool(pii) and not pol.allow_pii, "accessible": ok, "reason": reason,
            "documented": t.documented, **extra,
        })
    for item in _visible(ident, state.store.list()):
        try:
            s = state.store.get(item["id"])
        except KeyError:
            continue
        for d in s.datasets:
            v = state.services.validator.validate(d.sql, pol)
            doms = _report_domains(_source_tables([d.model_dump()]), kinds)
            objects.append({"type": "dataset", "id": f"{s.id}.{d.id}", "name": d.id, "business_name": "", "description": d.description,
                            "subject_area": doms[0] if doms else "Diğer", "report_id": s.id, "report_title": s.title,
                            "fields": len(d.fields), "view": d.view, "tables": v.tables, "accessible": v.ok,
                            "reason": None if v.ok else "; ".join(v.errors)[:200]})
    objects.sort(key=lambda o: (o["subject_area"], o["name"].lower()))
    return {"user": ident.public(), "role": ident.role,
            "policy": {"allowed_schemas": pol.allowed_schemas, "denied_tables": pol.denied_tables,
                       "allow_pii": pol.allow_pii, "max_rows": pol.max_rows},
            "catalog": {"ok": dd.catalog_error is None, "error": dd.catalog_error,
                        "undocumented": sum(1 for t in dd.tables.values() if not t.documented),
                        "missing": dd.missing_in_db[:30], "missing_count": len(dd.missing_in_db),
                        "renamed_count": len(dd.name_changes)},
            "objects": objects}



# --------------------------------------------------------------------------- yeni rapor önerileri
def _suggestion_items(request: Request, offset: int) -> list[dict[str, Any]]:
    """Kullanıcının yetkili olduğu veriden kural tabanlı öneriler (+ envanterde benzer rapor)."""
    from app import suggestions as sugg
    from app.harness.session import _source_tables
    ident = _need(request, "builder")
    pol = state.services.policy(ident.role)
    reports = []
    for item in _visible(ident, state.store.list()):
        try:
            s = state.store.get(item["id"])
        except KeyError:
            continue
        tables = _source_tables([d.model_dump() for d in s.datasets])
        if tables:
            reports.append({"id": s.id, "title": s.title, "tables": tables})
    return sugg.rule_suggestions(state.services.dictionary, lambda t: _object_access(pol, t)[0], reports, offset)


def _apply_texts(items: list[dict[str, Any]], texts: list[str] | None) -> bool:
    if not texts or len(texts) != len(items):
        return False
    for i, t in zip(items, texts):
        i["rule_text"], i["text"] = i["text"], t
    return True


@app.get("/api/suggestions")
def report_suggestions(request: Request, offset: int = 0) -> dict[str, Any]:
    """Yeni rapor önerileri — 1. katman (kural) anında; LLM ile düzenlenmiş hali önbellekte varsa o."""
    from app import suggestions as sugg
    items = _suggestion_items(request, offset)
    polished = _apply_texts(items, sugg.cached(get_settings().cache_dir, sugg.cache_key(state.gateway.s.llm_model, items)))
    return {"items": items, "polished": polished or not items, "offset": offset}


@app.post("/api/suggestions/polish")
def polish_suggestions(request: Request, offset: int = 0) -> dict[str, Any]:
    """2. katman: öneri iskeletlerini LLM ile iş diline çevirir (önbelleğe alınır). LLM yoksa kural metinleri döner."""
    from app import suggestions as sugg
    items = _suggestion_items(request, offset)
    polished = _apply_texts(items, sugg.polish(state.gateway, items, get_settings().cache_dir))
    return {"items": items, "polished": polished, "offset": offset}


# --------------------------------------------------------------------------- sorgu çalıştır (salt-okunur konsol)
QUERY_MAX_ROWS = 1000


@app.get("/api/query/schema")
def query_schema(request: Request) -> dict[str, Any]:
    """Sorgu ekranı için yetkili nesneler (ağaç + otomatik tamamlama): tablolar, onaylı view'lar, kolonlar, rapor dataset'leri."""
    ident = _need(request, "builder")
    pol = state.services.policy(ident.role)
    dd = state.services.dictionary
    kinds = dd._table_kinds()
    from app.harness.session import _source_tables
    view_sources = {str(v.get("name", "")).lower(): _source_tables([{"sql": v.get("original_sql")}])
                    for v in ViewRegistry(get_settings().views_registry).all() if v.get("original_sql")}
    objects = []
    for t in dd.tables.values():
        if not _object_access(pol, t)[0]:  # yalnız yetkili nesneler (politika + veritabanı SELECT yetkisi)
            continue
        kind = "view" if dd.is_view(t.name, kinds) else (kinds.get(t.name) if kinds.get(t.name) != "view" else None) or "table"
        # view'ın domain'i: kaynak SQL'indeki tabloların konu alanı (sözlükteki "Onaylı rapor view'ları" değil)
        area = t.subject_area
        if kind == "view":
            src = view_sources.get(t.name)
            area = ((_report_domains(src, kinds) if src else []) or _report_domains([t.name], kinds) or [t.subject_area])[0]
        objects.append({
            "id": t.name, "name": t.display_name or t.name, "kind": kind,
            "business_name": t.business_name, "description": t.description, "subject_area": area or "Diğer",
            "row_count": t.row_count, "documented": t.documented,
            "columns": [{"name": c.display_name or c.name, "type": c.data_type, "business_name": c.business_name,
                         "description": c.description, "pii": c.is_pii, "blocked": c.is_pii and not pol.allow_pii}
                        for c in t.columns],
        })
    objects.sort(key=lambda o: (o["kind"] == "view", o["subject_area"], o["name"].lower()))
    datasets = []
    for item in _visible(ident, state.store.list()):
        try:
            s = state.store.get(item["id"])
        except KeyError:
            continue
        for d in s.datasets:
            if state.services.validator.validate(d.sql, pol).ok:
                doms = _report_domains(_source_tables([d.model_dump()]), kinds)
                datasets.append({"report_id": s.id, "report_title": s.title, "id": d.id,
                                 "description": d.description, "sql": d.sql, "view": d.view,
                                 "subject_area": doms[0] if doms else "Diğer", "domains": doms})
    return {"role": ident.role, "max_rows": min(QUERY_MAX_ROWS, pol.max_rows), "allow_pii": pol.allow_pii,
            "allowed_schemas": pol.allowed_schemas, "objects": objects, "datasets": datasets}


class QueryIn(BaseModel):
    sql: str


@app.post("/api/query")
def run_query(body: QueryIn, request: Request) -> dict[str, Any]:
    """Kullanıcının yazdığı sorguyu rol yetkisi dahilinde, salt-okunur çalıştırır (en çok 1000 satır).
    Doğrulayıcı: tek SELECT, yetkili şema/tablo, PII, yasaklı fonksiyon; bağlantı readonly ve rollback."""
    ident = _need(request, "builder")
    pol = state.services.policy(ident.role)
    sql = (body.sql or "").strip()
    if len(sql) > 20000:
        raise HTTPException(413, "Sorgu çok uzun (en çok 20.000 karakter).")
    v = state.services.validator.validate(sql, pol, strict_joins=False, autofix=False)
    audit = Audit(get_settings().audit_log)
    limit = min(QUERY_MAX_ROWS, pol.max_rows)
    if not v.ok:
        audit.write(event="query_console", user=ident.username, role=ident.role, sql=sql[:4000], ok=False, errors=v.errors)
        return {"ok": False, "errors": v.errors, "warnings": v.warnings}
    try:
        res = state.services.connector.execute(v.sql, limit)
    except QueryError as e:
        msg = _short_db_error(str(e))
        audit.write(event="query_console", user=ident.username, role=ident.role, sql=sql[:4000], ok=False, errors=[msg])
        return {"ok": False, "errors": [f"Veritabanı hatası: {msg}"], "warnings": v.warnings}
    audit.write(event="query_console", user=ident.username, role=ident.role, sql=sql[:4000], ok=True,
                rows=len(res.rows), truncated=res.truncated, elapsed_ms=res.elapsed_ms, tables=v.tables)
    return {"ok": True, "columns": res.columns, "types": res.types, "rows": res.rows, "truncated": res.truncated,
            "row_limit": limit, "elapsed_ms": res.elapsed_ms, "tables": v.tables, "warnings": v.warnings}


@app.get("/api/sessions/{sid}")
def get_session(sid: str, request: Request) -> dict[str, Any]:
    return _own_session(request, sid)[0].public()


@app.delete("/api/sessions/{sid}")
def delete_session(sid: str, request: Request) -> dict[str, Any]:
    s, ident = _own_session(request, sid)
    rep = state.meta.report_by_session(sid) if state.meta is not None else None
    if rep and rep["status"] == "active":
        raise HTTPException(409, "Bu rapor Vitrin'de yayında. Silmeden önce yayından kaldırın.")
    state.store.delete(sid)
    _audit(ident, "session_delete", rep["report_id"] if rep else None, session=sid, title=s.title)
    return {"ok": True}


class MessageIn(BaseModel):
    content: str = ""
    images: list[str] | None = None


def _sse(events: Iterator[Event]) -> Iterator[str]:
    for ev in events:
        yield f"event: {ev.type}\ndata: {json.dumps(ev.data, ensure_ascii=False, default=str)}\n\n"


@app.post("/api/sessions/{sid}/messages")
def post_message(sid: str, body: MessageIn, request: Request) -> StreamingResponse:
    _s, ident = _own_session(request, sid)
    if not body.content.strip() and not body.images:
        raise HTTPException(400, "Mesaj boş")
    images = [i for i in (body.images or []) if i.startswith("data:image/")][:3]
    return StreamingResponse(_sse(state.agent.run_turn(sid, body.content, images, user=ident.username, role=ident.role)),
                             media_type="text/event-stream",
                             headers={"Cache-Control": "no-cache", "X-Accel-Buffering": "no"})


class TitleIn(BaseModel):
    title: str
    overwrite: bool = False   # aynı isimde başka rapor varsa üstüne yaz (o rapor silinir)


@app.put("/api/sessions/{sid}/title")
def rename_session(sid: str, body: TitleIn, request: Request) -> Any:
    _own_session(request, sid)
    title = " ".join(body.title.split())
    if not title:
        raise HTTPException(400, "Rapor adı boş olamaz.")
    if len(title) > 120:
        raise HTTPException(400, "Rapor adı en fazla 120 karakter olabilir.")
    with state.store.lock(sid):
        s = _session(sid)
        other = state.store.title_taken(title, sid, s.owner)
        if other and not body.overwrite:
            return JSONResponse({"detail": f"'{title}' adında başka bir rapor var.", "conflict_id": other}, status_code=409)
        s.title = title
        s.title_locked = True  # agent artık adı değiştirmez
        if s.spec:
            s.spec.title = title
            s.spec_version += 1
        if s.requirements:
            s.requirements.report_title = title
        state.store.save(s)   # aynı isimli diğer rapor (overwrite onaylandıysa) burada silinir
        return s.public()


class StatusIn(BaseModel):
    status: str


@app.put("/api/sessions/{sid}/status")
def set_status(sid: str, body: StatusIn, request: Request) -> dict[str, Any]:
    """Rapor yaşam döngüsü: idea (fikir) → design (tasarım) → test → live (canlıda).
    'live' = Vitrin'de yayında: statü ancak Yayınla ile canlıya alınır; canlıdan geri alınca yayın kaldırılır."""
    if body.status not in ("idea", "design", "test", "live"):
        raise HTTPException(400, "Geçersiz statü. Seçenekler: idea, design, test, live")
    _s, ident = _own_session(request, sid)
    rep = state.meta.report_by_session(sid) if state.meta is not None else None
    if body.status == "live" and state.meta is not None and not (rep and rep["status"] == "active"):
        raise HTTPException(409, "Raporu canlıya almak için Yayınla'yı kullanın (Vitrin'e sürüm ve paylaşım ile).")
    with state.store.lock(sid):
        s = _session(sid)
        s.status = body.status  # type: ignore[assignment]
        s.add(TranscriptItem(role="system", content=f"Rapor statüsü değişti: {body.status} ({ident.display_name})"))
        state.store.save(s)
    if rep and rep["status"] == "active" and body.status != "live":
        state.meta.set_status(rep["report_id"], "retired")
        _audit(ident, "report_retire", rep["report_id"], reason=f"statü {body.status}")
    _audit(ident, "status_change", rep["report_id"] if rep else None, session=sid, status=body.status)
    return s.public()


class PhaseIn(BaseModel):
    phase: str


@app.post("/api/sessions/{sid}/phase")
def set_phase(sid: str, body: PhaseIn, request: Request) -> dict[str, Any]:
    _own_session(request, sid)
    if body.phase not in PHASES:
        raise HTTPException(400, f"Geçersiz faz. Seçenekler: {PHASES}")
    with state.store.lock(sid):
        s = _session(sid)
        if body.phase == "data" and not s.requirements:
            raise HTTPException(409, "Önce gereksinimler kaydedilmeli.")
        if body.phase == "design" and not s.datasets:
            raise HTTPException(409, "Önce dataset'ler kaydedilmeli.")
        forward = PHASES.index(body.phase) > PHASES.index(s.phase)
        s.set_phase(body.phase)  # type: ignore[arg-type]

        if forward:
            s.add(TranscriptItem(role="system", content=f"Kullanıcı '{body.phase}' fazına geçti (kayıtlı içerikle)"))
            s.llm_messages.append({"role": "user", "content": "[HARNESS] Kullanıcı kayıtlı içerikle bu faza ileri geçti. "
                                   "Kayıtlı dataset'ler/dashboard ile devam et; durumu kısaca özetle ve ne yapmak istediğini sor."})
        else:
            s.add(TranscriptItem(role="system", content=f"Kullanıcı '{body.phase}' fazına döndü"))
            s.llm_messages.append({"role": "user", "content": "[HARNESS] Kullanıcı bu faza geri döndü. Mevcut durumu kısaca özetle ve ne değiştirmek istediğini sor."})
        state.store.save(s)
        return s.public()


# --------------------------------------------------------------------------- sorgu modu (hazır SQL ile rapor)
QUERY_PREVIEW_ROWS = 200
MAX_QUERY_DRAFTS = 20


class QueryDraftIn(BaseModel):
    id: str = ""
    title: str = ""
    sql: str = ""


class QueryDraftsIn(BaseModel):
    mode: str | None = None                       # "chat" | "query"
    drafts: list[QueryDraftIn] | None = None


@app.put("/api/sessions/{sid}/query-drafts")
def put_query_drafts(sid: str, body: QueryDraftsIn, request: Request) -> dict[str, Any]:
    """Sorgu modu ve taslak sorgular (henüz dataset değil); sayfa yenilense de kaybolmasın."""
    _own_session(request, sid)
    if body.mode is not None and body.mode not in ("chat", "query"):
        raise HTTPException(400, "Geçersiz mod (chat | query).")
    if body.drafts is not None:
        if len(body.drafts) > MAX_QUERY_DRAFTS:
            raise HTTPException(400, f"En çok {MAX_QUERY_DRAFTS} sorgu.")
        if any(len(d.sql) > 20000 for d in body.drafts):
            raise HTTPException(413, "Sorgu çok uzun (en çok 20.000 karakter).")
    with state.store.lock(sid):
        s = _session(sid)
        if body.mode:
            s.data_mode = body.mode  # type: ignore[assignment]
        if body.drafts is not None:
            s.query_drafts = [QueryDraft(**d.model_dump()) for d in body.drafts]
        state.store.save(s)
        return s.public()


@app.post("/api/sessions/{sid}/query-preview")
def query_preview(sid: str, body: QueryIn, request: Request) -> dict[str, Any]:
    """Sorgu modu önizlemesi: dataset kaydı ve dashboard ile AYNI doğrulama (önizlemede çalışan sorgu dashboard'da da çalışır).
    Rol, oturumun açıldığı andaki değil kullanıcının güncel rolüdür."""
    s, ident = _own_session(request, sid)
    sql = (body.sql or "").strip()
    if len(sql) > 20000:
        raise HTTPException(413, "Sorgu çok uzun (en çok 20.000 karakter).")
    ctx = ToolContext(s.model_copy(update={"user_role": ident.role}), state.services)   # oturum nesnesi değiştirilmez
    res, errors, tables, warnings = _run_validated(ctx, sql, QUERY_PREVIEW_ROWS)
    Audit(get_settings().audit_log).write(event="query_preview", session=sid, user=ident.username, role=ident.role,
                                          sql=sql[:4000], ok=not errors, errors=errors or None, tables=tables)
    if errors or res is None:
        return {"ok": False, "errors": errors, "warnings": warnings, "tables": tables}
    return {"ok": True, "columns": res.columns, "types": res.types, "rows": res.rows, "truncated": res.truncated,
            "row_limit": QUERY_PREVIEW_ROWS, "elapsed_ms": res.elapsed_ms, "tables": tables, "warnings": warnings}


class FromQueryIn(BaseModel):
    datasets: list[QueryDraftIn]


@app.post("/api/sessions/{sid}/datasets/from-query")
def datasets_from_query(sid: str, body: FromQueryIn, request: Request) -> Any:
    """Sorgu modundaki hazır SQL'leri dataset olarak kaydeder (agent'ın save_datasets'iyle aynı doğrulama ve profil).
    İhtiyaç / Veri fazındaysa Tasarım fazına geçilir ve agent'a veriyi özetleyip tasarımı sorma talimatı verilir;
    zaten tasarımdaysa dataset'ler eklenir / güncellenir (görseller korunur)."""
    _, ident = _own_session(request, sid)
    items = [d for d in body.datasets if d.sql.strip()]
    if not items:
        raise HTTPException(400, "En az bir sorgu gerekli.")
    ids = [d.id.strip() for d in items]
    if len(set(ids)) != len(ids):
        return JSONResponse({"detail": ["Sorgu adları benzersiz olmalı."]}, status_code=422)
    with state.store.lock(sid):
        s = _session(sid)
        s.user_role = ident.role   # güncel rol (rolü geri alınan kullanıcı eski yetkiyle kaydedemesin)
        ctx = ToolContext(s, state.services)
        built, errors = [], []
        for d in items:
            ds, prof, errs = _build_dataset(ctx, {"id": d.id.strip(), "sql": d.sql, "description": d.title.strip()}, user_sql=True)
            if errs:
                errors += errs
            else:
                built.append((ds, prof))
        if errors:
            return JSONResponse({"detail": errors}, status_code=422)
        for ds, prof in built:
            s.datasets = [x for x in s.datasets if x.id != ds.id] + [ds]
            s.dataset_profiles[ds.id] = prof
            if s.spec:
                s.spec.datasets = [x for x in s.spec.datasets if x.id != ds.id] + [ds]
        if s.spec:
            s.spec_version += 1
        if not s.requirements:
            title = s.title if s.title not in ("Yeni rapor", "") else (items[0].title.strip() or "Sorgu raporu")
            s.requirements = Requirements(report_title=title, business_goal="Kullanıcının hazır SQL sorgularıyla hazırlanan rapor",
                                          notes="Gereksinim sohbeti yapılmadı; veri kullanıcının sorgularından geldi.")
            if not s.title_locked and s.title in ("Yeni rapor", "") and title != "Sorgu raporu":
                s.title = title
        add_join_relationships(ctx)
        s.query_drafts = [QueryDraft(**d.model_dump()) for d in body.datasets]
        names = ", ".join(ds.id for ds, _ in built)
        tables = sorted({t for ds, _ in built for t in state.services.validator.validate(ds.sql, state.services.policy(ident.role)).tables})
        if s.phase != "design":
            s.set_phase("design")
            s.add(TranscriptItem(role="system", content=f"Sorgu modundan {len(built)} veri kümesi kaydedildi ({names}); tasarım fazına geçildi."))
            s.llm_messages.append({"role": "user", "content": f"[HARNESS] Kullanıcı ihtiyaç sohbeti yapmadan kendi hazır SQL sorgularıyla "
                                   f"{len(built)} dataset kaydetti ({names}). Ayrıntılı gereksinim kaydı yok; raporun amacını dataset'lerin kolonlarından ve "
                                   "profillerinden anla. " + _DESIGN_KICKOFF})
        else:
            s.add(TranscriptItem(role="system", content=f"Sorgu modundan veri kümesi eklendi / güncellendi: {names}."))
            s.llm_messages.append({"role": "user", "content": f"[HARNESS] Kullanıcı sorgu modundan dataset ekledi / güncelledi: {names}. "
                                   "Kısaca özetle ve bu veriyle ne eklemek istediğini sor."})
        Audit(get_settings().audit_log).write(event="datasets_from_query", session=sid, user=ident.username, role=ident.role,
                                              datasets=[ds.id for ds, _ in built], tables=tables)
        state.store.save(s)
        return s.public()


@app.put("/api/sessions/{sid}/spec")
def put_spec(sid: str, request: Request, spec: dict[str, Any] = Body(...)) -> Any:
    _own_session(request, sid)
    with state.store.lock(sid):
        s = _session(sid)
        # Elle düzenlemede SQL değişmiş/yeni dataset'ler de doğrulanır.
        for raw in spec.get("datasets") or []:
            known = next((d for d in s.datasets if d.id == raw.get("id")), None)
            if not known or known.sql.strip() != str(raw.get("sql", "")).strip():
                ds, prof, errs = _build_dataset(ToolContext(s, state.services), raw)
                if errs:
                    return JSONResponse({"detail": errs}, status_code=422)
                s.datasets = [d for d in s.datasets if d.id != ds.id] + [ds]
                s.dataset_profiles[ds.id] = prof
        new_spec, errs, _ = _validate_spec(ToolContext(s, state.services), spec)
        if errs:
            return JSONResponse({"detail": errs}, status_code=422)
        s.set_spec(new_spec)
        state.store.save(s)
        return s.public()


@app.post("/api/sessions/{sid}/demo")
def load_demo(sid: str, request: Request) -> Any:
    _own_session(request, sid)
    with state.store.lock(sid):
        s = _session(sid)
        raw = json.loads(get_settings().demo_spec.read_text(encoding="utf-8"))
        ctx = ToolContext(s, state.services)
        built, profiles = [], {}
        for d in raw["datasets"]:
            ds, prof, errs = _build_dataset(ctx, d)
            if errs:
                raise HTTPException(500, f"Demo dataset hatası: {errs}")
            # demo'daki elle verilmiş etiket/formatları koru
            ds.fields = [DatasetField.model_validate(f) for f in d.get("fields") or []] or ds.fields
            built.append(ds)
            profiles[ds.id] = prof
        s.datasets, s.dataset_profiles = built, profiles
        req = raw.pop("requirements", None) or {}
        s.requirements = Requirements.model_validate({
            "report_title": raw["title"], "business_goal": raw.get("subtitle") or "", "audience": "Yönetim",
            "kpis": [v["title"] for v in raw["visuals"] if v["type"] == "kpi"],
            "dimensions": [], **req})
        spec = ReportSpec.model_validate({**raw, "datasets": [d.model_dump() for d in built]})
        s.set_phase("design")
        # Aynı başlıklı rapor kaydedilince eskisi silinir (üstüne yazma kuralı): ikinci kez yüklenen demo, ilk demo
        # raporunu (Vitrin'de yayında olsa bile) sessizce siliyordu. Başlık doluysa numaralanır.
        title, n = raw["title"], 2
        while not s.title_locked and state.store.title_taken(title, sid, s.owner):
            title, n = f"{raw['title']} ({n})", n + 1
        if s.title_locked:
            title = s.title
        spec.title = title
        s.set_spec(spec)
        s.title = title

        s.add(TranscriptItem(role="system", content="Demo dashboard yüklendi. Değişiklik isteyebilirsiniz (ör. \"bölge grafiğini donut yap\")."))
        s.llm_messages = [{"role": "user", "content": "[HARNESS] Demo dashboard yüklendi. Kullanıcının değişiklik isteklerini bekle."}]
        state.store.save(s)
        return s.public()


# --------------------------------------------------------------------------- dashboard verisi
_cache: dict[str, tuple[float, dict[str, Any]]] = {}
_cache_lock = threading.Lock()
CACHE_TTL_S = 120


def _dataset_payload(sql: str, role: str) -> dict[str, Any]:
    key = hashlib.sha1(f"{role}|{sql}".encode()).hexdigest()
    with _cache_lock:
        hit = _cache.get(key)
        if hit and time.time() - hit[0] < CACHE_TTL_S:
            return hit[1]
    pol = state.services.policy(role)
    v = state.services.validator.validate(sql, pol)
    if not v.ok:   # rolün izin vermediği veri (şema / tablo / kişisel veri): görsel "yetkiniz yok" gösterir
        return {"columns": [], "rows": [], "error": "; ".join(v.errors), "denied": True}
    try:
        r = state.services.connector.execute(v.sql, pol.max_rows)
    except QueryError as e:
        return {"columns": [], "rows": [], "error": str(e)[:300]}
    payload = {"columns": r.columns, "rows": r.rows, **({"truncated": True} if r.truncated else {})}
    with _cache_lock:
        _cache[key] = (time.time(), payload)
        if len(_cache) > 500:
            _cache.clear()
    return payload


def _all_datasets(s) -> dict[str, Any]:
    datasets = {d.id: d for d in s.datasets}
    if s.spec:
        datasets.update({d.id: d for d in s.spec.datasets})
    return datasets


def _engine(s=None) -> ModelFilterEngine:
    """Filtre motoru; rapor verilirse o rapora özel (onaylı) ilişkiler de kullanılır."""
    from app.dictionary.model_rels import report_relationships
    extra = report_relationships(state.services.dictionary, getattr(s, "model_relationships", None)) if s is not None else []
    return ModelFilterEngine(state.services.dictionary, state.services.connector.dialect, extra)


def _bindings(s) -> dict[str, dict[str, str]]:
    """dataset → {alan: 'şema.tablo.kolon'}: hangi dataset kolonu hangi model kolonundan geliyor (tıklayarak filtre için)."""
    eng = _engine(s)
    return {did: eng.lineage(d.sql) for did, d in _all_datasets(s).items()}


def _filter_defs(s, bindings: dict[str, dict[str, str]] | None = None) -> list[dict[str, Any]]:
    """Spec filtrelerini model kolonuna çözer (table+column ya da field'ın kökeni)."""
    if not s.spec:
        return []
    out = []
    for f in s.spec.filters:
        key = f"{f.table}.{f.column}".lower() if f.table and f.column else None
        if not key and f.field:
            bindings = bindings if bindings is not None else _bindings(s)
            key = next((b[f.field] for b in bindings.values() if f.field in b), None)
        out.append({"id": f.id, "label": f.label, "type": f.type, "field": f.field, "key": key})
    return out


_options_cache: dict[str, tuple[float, list[Any]]] = {}


def _filter_options(key: str, role: str) -> list[Any]:
    hit = _options_cache.get(f"{role}|{key}")
    if hit and time.time() - hit[0] < CACHE_TTL_S * 5:
        return hit[1]
    table, _, column = key.rpartition(".")
    dd = state.services.dictionary
    t = dd.tables.get(table)
    col = next((c for c in t.columns if c.name == column), None) if t else None
    if not t or not col:
        return []
    dialect = state.services.connector.dialect
    c = exp.column(col.display_name or column)
    where = c.is_(exp.null()).not_()
    if t.snapshot_date:   # günlük anlık görüntü: seçenekler son günden (takvimle çoğaltılmış view'ın tamamı taranmaz)
        d = exp.column(t.snapshot_date)
        from app.data.snapshot_guard import recent_bound
        last = exp.select(exp.Max(this=d.copy())) \
            .from_(with_nolock(exp.to_table(t.display_name or table, dialect=dialect), dd)).where(recent_bound(d.copy()))
        where = exp.and_(where, exp.EQ(this=d.copy(), expression=exp.Subquery(this=last)))
    sql = exp.select(c).distinct().from_(with_nolock(exp.to_table(t.display_name or table, dialect=dialect), dd)).where(where)         .order_by(c).limit(500).sql(dialect=dialect)
    payload = _dataset_payload(sql, role)  # validator: izinli şema + PII kontrolü burada da geçerli
    values = [r[0] for r in payload.get("rows", [])]
    if payload.get("error"):   # hata (ör. zaman aşımı) önbelleğe alınmaz: boş liste 10 dk takılı kalmasın
        log.warning("Filtre seçenekleri alınamadı (%s): %s", key, str(payload["error"])[:200])
        return values
    _options_cache[f"{role}|{key}"] = (time.time(), values)
    return values


AS_OF_KEY = "__as_of__"   # dashboard seçimi: günlük anlık görüntülerde 'itibarıyla' tarihi (YYYY-MM-DD)
_DAY = re.compile(r"^\d{4}-\d{2}-\d{2}$")


def _dashboard_data(s, selections: list[dict[str, Any]] | None = None) -> dict[str, Any]:
    """selections: [{key: 'şema.tablo.kolon', values: [...], exclude: [dataset_id, ...]}];
    key '__as_of__' → günlük anlık görüntülerin veri tarihi (boşsa son gün)."""
    eng = _engine(s)
    out: dict[str, Any] = {}
    applied: dict[str, list[str]] = {}
    as_of = next((str(sel["values"][0])[:10] for sel in selections or []
                  if sel.get("key") == AS_OF_KEY and sel.get("values")), None)
    # filtreler varsayılan olarak tüm görselleri etkiler; yalnız TÜM görselleri ignoreFilters olan dataset'ler muaf
    exempt: set[str] = set()
    if s.spec:
        by_ds: dict[str, list[bool]] = {}
        for v in s.spec.visuals:
            if v.datasetId:
                by_ds.setdefault(v.datasetId, []).append(bool(v.options.ignoreFilters))
        exempt = {d for d, flags in by_ds.items() if flags and all(flags)}
    if as_of and not _DAY.match(as_of):
        raise HTTPException(400, "Veri tarihi YYYY-AA-GG biçiminde olmalı.")
    for did, d in _all_datasets(s).items():
        sql = d.sql
        dated = False
        if as_of:
            try:
                sql, dated = eng.as_of(d.sql, as_of)
                if not dated and d.view and d.original_sql:   # view toplulaştırılmış: tarih kaynak SQL'e uygulanır
                    alt, dated = eng.as_of(d.original_sql, as_of)
                    sql = alt if dated else sql
            except Exception as e:  # noqa: BLE001
                log.warning("Veri tarihi uygulanamadı (%s): %s", did, e)
        base_sql = sql
        active = [ModelFilter(*sel["key"].rsplit(".", 1), list(sel["values"]))
                  for sel in selections or [] if sel.get("values") and did not in (sel.get("exclude") or [])
                  and "." in sel.get("key", "") and sel.get("key") != AS_OF_KEY and did not in exempt]
        if active:
            try:
                sql, applied[did] = eng.apply(base_sql, active)
            except Exception as e:  # noqa: BLE001 — filtre uygulanamazsa filtresiz göster
                log.warning("Filtre uygulanamadı (%s): %s", did, e)
                applied[did] = []
            # onaylı view'a bağlı dataset: view toplulaştırılmış olduğundan (ör. yalnız bölge) filtre ona ulaşamayabilir.
            # Bu durumda view'ın kaynak SQL'i filtrelenerek çalıştırılır (kolonlar aynı; filtresizken view kullanılır).
            if d.view and d.original_sql and len(applied.get(did, [])) < len({f.key for f in active}):
                try:
                    src = eng.as_of(d.original_sql, as_of)[0] if as_of else d.original_sql
                    alt_sql, alt_applied = eng.apply(src, active)
                    if len(alt_applied) > len(applied.get(did, [])):
                        sql, applied[did] = alt_sql, alt_applied
                except Exception as e:  # noqa: BLE001
                    log.warning("View kaynağına filtre uygulanamadı (%s): %s", did, e)
        if dated:
            applied.setdefault(did, []).append(AS_OF_KEY)
        out[did] = _dataset_payload(sql, s.user_role)
    return {"datasets": out, "applied": applied, **({"as_of": as_of} if as_of else {})}


def _data_date_info(s, role: str) -> dict[str, Any] | None:
    """Rapor günlük anlık görüntü okuyorsa veri tarihi seçicisi: kolon, tablolar, en eski / en yeni gün."""
    eng = _engine(s)
    tables: list[str] = []
    for d in _all_datasets(s).values():
        for sql in (d.sql, d.original_sql):
            if sql:
                tables += eng.snapshot_tables(sql)
    if not tables:
        return None
    main = max(set(tables), key=tables.count)
    dd = state.services.dictionary
    t = dd.tables[main]
    key = f"{role}|asof|{main}"
    hit = _options_cache.get(key)
    if hit and time.time() - hit[0] < CACHE_TTL_S * 5:
        lo, hi = hit[1]
    else:
        dialect = state.services.connector.dialect
        # yalnız son gün okunur (T-1 veri ambarı: son 2 gün; yoksa 31 gün) — tüm geçmişin MIN'i taranmaz,
        # seçicinin alt sınırı boş kalır (kullanıcı geçmiş bir günü yine seçebilir)
        from app.data.snapshot_guard import recent_bound
        lo, hi = None, None
        for days in (2, 31):
            c = exp.column(t.snapshot_date)
            sql = exp.select(exp.Max(this=c).as_("max_d")) \
                .from_(with_nolock(exp.to_table(t.display_name or main, dialect=dialect), dd)) \
                .where(recent_bound(c.copy(), days)).sql(dialect=dialect)
            rows = _dataset_payload(sql, role).get("rows") or []
            hi = str(rows[0][0])[:10] if rows and rows[0][0] else None
            if hi:
                break
        if lo or hi:   # boş / hatalı sonuç önbelleğe alınmaz
            _options_cache[key] = (time.time(), [lo, hi])
    return {"key": AS_OF_KEY, "column": t.snapshot_date, "table": t.display_name or main,
            "tables": sorted({dd.tables[x].display_name or x for x in tables}), "min": lo, "max": hi}


@app.get("/api/sessions/{sid}/dashboard-data")
def dashboard_data(sid: str, request: Request) -> dict[str, Any]:
    return _dashboard_data(_own_session(request, sid)[0])


class SelectionIn(BaseModel):
    key: str
    values: list[Any]
    exclude: list[str] | None = None


class DashboardQuery(BaseModel):
    selections: list[SelectionIn] = []


@app.post("/api/sessions/{sid}/dashboard-data")
def dashboard_data_filtered(sid: str, body: DashboardQuery, request: Request) -> dict[str, Any]:
    """Model filtreleriyle dashboard verisi: seçimler ilişkiler üzerinden her dataset'e yayılır."""
    return _dashboard_data(_own_session(request, sid)[0], [sel.model_dump() for sel in body.selections])


@app.get("/api/sessions/{sid}/filters")
def dashboard_filters(sid: str, request: Request) -> dict[str, Any]:
    """Dilimleyici tanımları + seçenekleri ve dataset alanlarının model kökenleri."""
    s = _own_session(request, sid)[0]
    bindings = _bindings(s)
    defs = _filter_defs(s, bindings)
    for f in defs:
        f["options"] = _filter_options(f["key"], s.user_role) if f["key"] else []
    try:
        data_date = _data_date_info(s, s.user_role)
    except Exception as e:  # noqa: BLE001 — tarih seçicisi olmadan da dashboard çalışsın
        log.warning("Veri tarihi aralığı alınamadı: %s", e)
        data_date = None
    return {"filters": defs, "bindings": bindings, **({"data_date": data_date} if data_date else {})}


@app.get("/api/sessions/{sid}/export/html")
def export_html(sid: str, request: Request) -> HTMLResponse:
    s = _own_session(request, sid)[0]
    if not s.spec:
        raise HTTPException(409, "Henüz dashboard yok.")
    viewer = get_settings().viewer_html
    if not viewer.exists():
        raise HTTPException(503, "Viewer derlenmemiş: frontend klasöründe `npm run build:viewer` çalıştırın.")
    data = _dashboard_data(s)
    data["bindings"] = _bindings(s)
    data["filter_keys"] = {f["id"]: f["key"] for f in _filter_defs(s, data["bindings"])}
    payload = json.dumps({"spec": s.spec.model_dump(exclude_none=True), "data": data},
                         ensure_ascii=False, default=str).replace("<", "\\u003c")
    html = viewer.read_text(encoding="utf-8").replace("__REPORT_JSON__", payload, 1)
    title = s.spec.title or "dashboard"
    ascii_name = "".join(c if c.isascii() and c.isalnum() else "_" for c in title)[:60] or "dashboard"
    return HTMLResponse(html, headers={"Content-Disposition": f"attachment; filename=\"{ascii_name}.html\"; "
                                                              f"filename*=UTF-8''{quote(title[:60])}.html"})


# --------------------------------------------------------------------------- sözlük
# --------------------------------------------------------------------------- dataset → onaylı view
class ViewIn(BaseModel):
    name: str | None = None      # view adı (şemasız); boşsa v_<dataset_id>


def _find_dataset(s, did: str):
    ds = next((d for d in s.datasets if d.id == did), None) or (next((d for d in s.spec.datasets if d.id == did), None) if s.spec else None)
    if not ds:
        raise HTTPException(404, f"Dataset bulunamadı: {did}")
    return ds


@app.post("/api/sessions/{sid}/datasets/{did}/view-script")
def dataset_view_script(sid: str, did: str, body: ViewIn, request: Request) -> dict[str, Any]:
    """Dataset için inceleme amaçlı CREATE OR ALTER VIEW scripti üretir (çalıştırmaz)."""
    s = _own_session(request, sid)[0]
    ds = _find_dataset(s, did)
    settings = get_settings()
    name = (body.name or safe_view_name(did)).strip()
    source_sql = ds.original_sql or ds.sql
    v = state.services.validator.validate(source_sql, state.services.policy(s.user_role))
    if not v.ok:
        raise HTTPException(422, "; ".join(v.errors))
    vs = build_view_script(schema=settings.view_schema, name=name, dataset_id=did, description=ds.description,
                           validated_sql=v.sql, columns=[f.name for f in ds.fields], session_title=s.title)
    if vs.errors:
        return JSONResponse({"detail": vs.errors}, status_code=422)
    settings.view_scripts_dir.mkdir(parents=True, exist_ok=True)
    path = settings.view_scripts_dir / f"{settings.view_schema}.{name}.sql"
    path.write_text(vs.script, encoding="utf-8")
    exists = view_columns(state.services.connector, settings.view_schema, name) is not None
    return {"view": vs.name, "script": vs.script, "file": str(path), "exists": exists}


@app.post("/api/sessions/{sid}/datasets/{did}/use-view")
def dataset_use_view(sid: str, did: str, body: ViewIn, request: Request) -> Any:
    """View oluşturulduysa: doğrular, sözlüğe ekler, dataset'i view'dan okuyacak şekilde değiştirir."""
    _own_session(request, sid)
    settings = get_settings()
    name = (body.name or safe_view_name(did)).strip()
    with state.store.lock(sid):
        s = _session(sid)
        ds = _find_dataset(s, did)
        cols = view_columns(state.services.connector, settings.view_schema, name)
        if cols is None:
            raise HTTPException(409, f"{settings.view_schema}.{name} veritabanında bulunamadı. Scripti çalıştırdınız mı?")
        want = [f.name for f in ds.fields]
        missing = [c for c in want if c.lower() not in {x.lower() for x in cols}]
        if missing:
            raise HTTPException(409, f"View kolonları dataset ile uyuşmuyor; eksik: {missing}. Scripti yeniden üretip çalıştırın.")
        original = ds.original_sql or ds.sql
        eng = _engine(s)
        lineage = eng.lineage(original)
        entry = {"name": f"{settings.view_schema}.{name}", "business_name": ds.description or did,
                 "description": f"'{s.title}' raporunun '{did}' dataset'i (onaylı view).", "dataset_id": did,
                 "session_id": sid, "original_sql": original, "created_at": now_iso(),
                 "columns": [{"name": f.name, "label": f.label, "type": f.type, "format": f.format,
                              "lineage": lineage.get(f.name)} for f in ds.fields]}
        ViewRegistry(settings.views_registry).upsert(entry)
        state.services.dictionary.add_view(entry)
        # hangi filtreler view'a geçince artık uygulanamıyor?
        filters = [f for f in _filter_defs(s) if f["key"]]
        new_sql = select_from_view(settings.view_schema, name, want)
        def applied(sql: str) -> set[str]:
            try:
                return set(eng.apply(sql, [ModelFilter(*f["key"].rsplit(".", 1), ["__probe__"]) for f in filters])[1])
            except Exception:  # noqa: BLE001
                return set()
        lost = [f["label"] for f in filters if f["key"] in applied(original) - applied(new_sql)]
        for d in [*s.datasets, *(s.spec.datasets if s.spec else [])]:
            if d.id == did:
                d.original_sql, d.sql, d.view = original, new_sql, entry["name"]
        s.spec_version += 1
        s.add(TranscriptItem(role="system", content=f"'{did}' dataset'i artık onaylı view'dan okunuyor: {entry['name']}"
                             + (f" (bu görsele artık uygulanamayan filtreler: {', '.join(lost)})" if lost else "")))
        state.store.save(s)
        return {"session": s.public(), "view": entry["name"], "lost_filters": lost}


@app.get("/api/views")
def list_views(request: Request) -> list[dict[str, Any]]:
    _need(request, "builder")
    return ViewRegistry(get_settings().views_registry).all()


@app.get("/api/dictionary/search")
def dictionary_search(request: Request, q: str = "") -> list[dict[str, Any]]:
    _need(request, "builder")
    return state.services.dictionary.search(q, 10) if q.strip() else []


@app.get("/api/dictionary/model")
def dictionary_model(request: Request, session: str | None = None) -> dict[str, Any]:
    """İlişkisel model; session verilirse o rapordaki dataset'lerin kullandığı tablolar da işaretlenir."""
    _need(request, "builder")
    model = state.services.dictionary.model()
    used: dict[str, list[str]] = {}
    if session:
        s = _own_session(request, session)[0]
        datasets = {d.id: d for d in s.datasets}
        if s.spec:
            datasets.update({d.id: d for d in s.spec.datasets})
        pol = state.services.policy(s.user_role)
        for d in datasets.values():
            for t in state.services.validator.validate(d.sql, pol).tables:
                used.setdefault(t, []).append(d.id)
    if session and s.datasets and not s.joins_derived:
        # eski raporlar: dataset JOIN'lerinden rapora özel ilişkiler bir kez çıkarılır (Model sekmesi + filtreler)
        from app.harness.tools import add_join_relationships
        lock = state.store.lock(session)
        if lock.acquire(blocking=False):      # agent çalışıyorsa bekleme: sonraki açılışta çıkarılır
            try:
                s = _session(session)
                if not s.joins_derived:
                    add_join_relationships(ToolContext(s, state.services))
                    state.store.save(s)
            finally:
                lock.release()
    if session:
        dd = state.services.dictionary
        disp = lambda t, c: next((x.display_name for x in dd.tables[t].columns if x.name == c and x.display_name), c)  # noqa: E731
        from app.dictionary.model_rels import report_relationships
        for r in report_relationships(dd, s.model_relationships):
            model["relationships"].append({"id": r.id, "from_table": r.from_table, "to_table": r.to_table,
                                           "pairs": [[a, b] for a, b in r.pairs],
                                           "pairs_display": [[disp(r.from_table, a), disp(r.to_table, b)] for a, b in r.pairs],
                                           "cardinality": r.cardinality, "role": r.role, "active": True,
                                           "source": "sql" if r.id.startswith("join:") else "report",
                                           "description": r.description})
    model["used_tables"] = used
    model["dialect"] = state.services.connector.dialect
    return model


@app.delete("/api/dictionary/relationships/{rel_id:path}")
def delete_model_relationship(rel_id: str, request: Request) -> dict[str, Any]:
    """Kullanıcı onayıyla eklenmiş (agent'ın önerdiği) model ilişkisini kaldırır — yalnız admin."""
    from app.dictionary.model_rels import RelationshipRegistry
    ident = _need(request, "admin")
    if not rel_id.startswith("model:"):
        raise HTTPException(400, "Yalnız sonradan eklenen model ilişkileri silinebilir (sözlük / veritabanı ilişkileri değil).")
    removed = RelationshipRegistry(get_settings().model_relationships).remove(rel_id)
    dd = state.services.dictionary
    dd.relationships = [r for r in dd.relationships if r.id != rel_id]
    if removed:
        _audit(ident, "model_relationship_delete", None, relationship=rel_id)
    return {"ok": removed}


@app.post("/api/dictionary/reload")
def dictionary_reload(request: Request) -> dict[str, Any]:
    _need(request, "admin")
    state.services.dictionary.load()
    return {"ok": True, "tables": len(state.services.dictionary.tables)}


# --------------------------------------------------------------------------- bağlantı ayarları
from app.data import connections as conns  # noqa: E402
from app.data.odbc import friendly_error  # noqa: E402


def _require_settings_access(request: Request) -> Identity:
    """Bağlantı ayarları: admin rolü ya da uygulamanın çalıştığı bilgisayarın kendisi (ağdan gelen kullanıcı değil).
    Sunucu modunda yalnız platform yöneticisi (admin)."""
    ident = _me(request)
    if get_settings().platform_mode == "server":
        if not authz.at_least(ident, "admin"):
            raise HTTPException(403, "Bağlantı ayarlarını yalnızca platform yöneticisi (admin) değiştirebilir.")
        return ident
    fwd = [x.strip() for x in (request.headers.get("x-forwarded-for") or "").split(",") if x.strip()]
    local = (request.client.host if request.client else "") in ("127.0.0.1", "::1", "localhost", "testclient") \
        and all(x in ("127.0.0.1", "::1", "::ffff:127.0.0.1") for x in fwd)
    if ident.role != "admin" and not local:
        raise HTTPException(403, "Bağlantı ayarlarını yalnızca admin rolü ya da sunucunun çalıştığı bilgisayar değiştirebilir.")
    return ident


class ConnIn(BaseModel):
    server: str = "localhost"
    database: str = ""
    auth: str = "windows"            # windows | sql
    username: str = ""
    password: str | None = None      # None → kayıtlı şifre korunur
    encrypt: bool = True
    trust_server_certificate: bool = True
    same_as_data: bool = True        # yalnız sözlük için
    sources: dict[str, list[str]] | None = None   # yalnız sözlük: rol → tablolar / sayfalar
    kind: str = "sqlserver"          # yalnız sözlük: sqlserver | mysql | postgres | excel | none
    port: int | None = None          # MySQL (varsayılan 3306) / PostgreSQL (5432)
    excel_path: str = ""             # Excel dosyası (yüklenen ya da ağ yolu)
    extra_databases: list[str] = []  # yalnız veri kaynağı: aynı sunucudaki ek veritabanları (ör. EDW)
    mappings: dict[str, dict[str, str]] | None = None   # yalnız sözlük: kaynak → alan → başlık ("" = kullanma)


class ConnectionsIn(BaseModel):
    data: ConnIn
    dictionary: ConnIn


class ConnTestIn(BaseModel):
    target: str = "data"             # data | dictionary
    data: ConnIn
    dictionary: ConnIn | None = None


def _fields(c: ConnIn, saved: dict[str, Any] | None) -> dict[str, Any]:
    d = c.model_dump()
    if d.get("kind") in ("mysql", "postgres"):   # MySQL / PostgreSQL her zaman kullanıcı adı / şifre ile
        d["auth"] = "sql"
    if d["auth"] != "sql":
        d.update(username="", password="")
    elif d["password"] is None:      # boş bırakıldı: kayıtlı şifreyi kullan
        d["password"] = conns.unprotect((saved or {}).get("password_enc") or "") if saved else ""
    d["server"] = (d["server"] or "localhost").strip()
    d["database"] = (d["database"] or "").strip()
    d["extra_databases"] = [x for x in conns._db_list(d.get("extra_databases") or []) if x.lower() != d["database"].lower()]
    return d


def _resolved(body: ConnTestIn | ConnectionsIn) -> tuple[dict[str, Any], dict[str, Any]]:
    saved = conns.load_connections() or {}
    data = _fields(body.data, saved.get("data"))
    dic_in = body.dictionary or ConnIn(database=conns.DEFAULT_DICTIONARY_DB)
    from app.dictionary.sources import KINDS, ROLES, default_sources
    kind = dic_in.kind if dic_in.kind in KINDS else "sqlserver"
    dic = _fields(dic_in, saved.get("dictionary"))
    if kind == "sqlserver" and dic_in.same_as_data:
        dic = {**data, "database": dic["database"] or conns.DEFAULT_DICTIONARY_DB, "same_as_data": True}
    else:
        dic["same_as_data"] = False
    dic["kind"] = kind
    if kind in ("mysql", "postgres"):
        dic["port"] = dic_in.port or (3306 if kind == "mysql" else 5432)
    if kind == "excel":
        dic["excel_path"] = (dic_in.excel_path or "").strip().strip('"')
    src = dic_in.sources or default_sources(kind)
    dic["sources"] = {r: [n.strip() for n in (src.get(r) or []) if n and n.strip()] for r in ROLES}
    used = {n.casefold() for lst in dic["sources"].values() for n in lst}
    dic["mappings"] = {k: {f: str(h or "") for f, h in (v or {}).items()}
                       for k, v in (dic_in.mappings or {}).items() if k.casefold() in used}
    for k in ("sources", "kind", "port", "excel_path", "mappings"):
        data.pop(k, None)
    return data, dic


def _probe(fields: dict[str, Any], database: str | None = None) -> dict[str, Any]:
    from app.data.connector import SqlServerConnector
    odbc = conns.build_odbc(fields, database)
    con = SqlServerConnector(odbc, 15)
    r = con.execute("SELECT @@SERVERNAME, DB_NAME(), CAST(SERVERPROPERTY('ProductVersion') AS nvarchar(64)), "
                    "CAST(SERVERPROPERTY('Edition') AS nvarchar(128)), SUSER_SNAME()", 1)
    srv, db, ver, ed, login = r.rows[0]
    return {"server_name": srv, "database": db, "version": f"SQL Server {ver}", "edition": ed, "login": login,
            "driver": odbc.split(";")[0].replace("DRIVER=", "").strip("{}")}


@app.post("/api/settings/dictionary/upload")
async def upload_dictionary_excel(request: Request, filename: str = "sozluk.xlsx") -> dict[str, Any]:
    """Excel sözlük dosyasını yükler (gövde: dosyanın kendisi). Sayfalar ve önerilen roller döner."""
    _require_settings_access(request)
    import re as _re
    ext = Path(filename).suffix.lower()
    if ext not in (".xlsx", ".xlsm"):
        raise HTTPException(400, "Yalnızca .xlsx / .xlsm dosyaları yüklenebilir.")
    body = await request.body()
    if not body or len(body) > 30 * 1024 * 1024:
        raise HTTPException(400, "Dosya boş ya da 30 MB'tan büyük.")
    safe = _re.sub(r"[^\w.\-]+", "_", Path(filename).stem)[:60] or "sozluk"
    conns.DICTIONARY_FILES_DIR.mkdir(parents=True, exist_ok=True)
    path = conns.DICTIONARY_FILES_DIR / f"{safe}{ext}"
    path.write_bytes(body)
    from app.dictionary.sources import ExcelReader, _norm_key, suggest_role
    try:
        sheets = ExcelReader(path).list_tables()
    except Exception as e:  # noqa: BLE001
        path.unlink(missing_ok=True)
        raise HTTPException(400, f"Excel dosyası okunamadı: {e}")
    return {"ok": True, "path": str(path),
            "tables": [{"name": n, "columns": c, "role": suggest_role({_norm_key(x) for x in c})} for n, c in sheets.items()]}


@app.get("/api/settings/dictionary/template.xlsx")
def dictionary_template(request: Request, layout: str = "multi") -> Response:
    """Şu an yüklü sözlüğü Excel olarak indirir; düzenleyip Excel kaynağı olarak geri yüklenebilir.
    layout=multi: Tablolar / Kolonlar / İlişkiler / Metrikler sayfaları · layout=single: tek "Sözlük" sayfası
    (her satır bir kolon, tablo bilgileri table_* kolonlarında; ilişkiler foreign key / kolon adlarından çıkarılır)."""
    _require_settings_access(request)
    import io

    import openpyxl
    from openpyxl.styles import Font, PatternFill
    from app.dictionary.sources import DEFAULTS, ROLES
    dd = state.services.dictionary
    _disp = lambda t: (dd.tables[t].display_name or t) if t in dd.tables else t  # noqa: E731 — özgün yazım (dbo.FactX)
    rows: dict[str, list[list[Any]]] = {
        "tables": [[t.display_name or t.name, t.business_name, t.description, t.subject_area, t.grain, t.row_count, t.table_type]
                   for t in dd.tables.values() if t.table_type != "view"],
        "columns": [[t.display_name or t.name, c.display_name or c.name, c.business_name, c.description, c.data_type, c.role,
                     c.default_aggregation, ", ".join(c.synonyms), "Evet" if c.is_pii else "", c.sample_values]
                    for t in dd.tables.values() if t.table_type != "view" for c in t.columns],
        # onaylı view'lar (ve onların ilişkileri) sözlüğe uygulamaca eklenir; şablona girmez
        "relationships": [[r.id, _disp(r.from_table), a, _disp(r.to_table), b, r.cardinality, r.role, "Evet" if r.active else "Hayır"]
                          for r in dd.relationships for a, b in r.pairs
                          if not any((dd.tables.get(t) and dd.tables[t].table_type == "view") for t in (r.from_table, r.to_table))],
        "metrics": [[m.name, m.expression_sql, m.business_name, m.description, m.base_table, m.value_format, ", ".join(m.synonyms)]
                    for m in dd.metrics],
    }
    order = {"relationships": ["relationship_id", "from_table", "from_column", "to_table", "to_column", "cardinality", "role", "is_active"]}
    order["columns"] = ["table_name", "column_name", "business_name", "description", "data_type", "column_role",
                        "default_aggregation", "synonyms", "is_pii", "sample_values"]
    order["tables"] = ["table_name", "business_name", "description", "subject_area", "grain", "row_count", "table_type"]
    sheets: list[tuple[str, list[str], list[list[Any]]]]
    if layout == "single":
        tinfo = {r[0].lower(): r for r in rows["tables"]}
        flat = []
        for c in rows["columns"]:
            t = tinfo.get(str(c[0]).lower(), [c[0], None, None, None, None, None, None])
            flat.append([c[0], t[1], t[2], t[3], t[4], t[5], t[6], *c[1:]])
        sheets = [("Sözlük", ["table_name", "table_business_name", "table_description", "subject_area", "grain", "row_count",
                              "table_type", *order["columns"][1:]], flat)]
    else:
        sheets = [(DEFAULTS["excel"][role][0], order[role] if role in order else spec["required"] + spec["optional"], rows[role])
                  for role, spec in ROLES.items()]
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for title, header, data in sheets:
        ws = wb.create_sheet(title)
        ws.append(header)
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="2563EB")
        for r in data:
            ws.append(r)
        ws.freeze_panes = "A2"
        for i, h in enumerate(header, 1):
            ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = max(14, min(40, len(h) + 6))
    buf = io.BytesIO()
    wb.save(buf)
    return Response(buf.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f'attachment; filename="veri_sozlugu{"_tek_sayfa" if layout == "single" else ""}.xlsx"'})


# ---------------------------------------------------------------- dil modeli (LLM) bağlantısı
def _clean_key(v: str | None) -> str | None:
    """Yapıştırırken gelen boşluk / satır sonu / görünmez karakterleri atar (HTTP başlığını bozup 'bağlanılamadı' verir)."""
    return None if v is None else "".join(ch for ch in v if ch.isprintable() and not ch.isspace())


def _clean_url(v: str | None) -> str:
    v = (v or "").strip().rstrip("/")
    if v and "://" not in v:
        v = "http://" + v
    return v


class VisionIn(BaseModel):
    enabled: bool = True
    same_as_main: bool = True
    base_url: str = ""
    api_key: str | None = None       # None → kayıtlı anahtar korunur
    model: str = ""

    @field_validator("api_key")
    @classmethod
    def _k(cls, v): return _clean_key(v)

    @field_validator("base_url")
    @classmethod
    def _u(cls, v): return _clean_url(v)


class LlmIn(BaseModel):
    base_url: str
    api_key: str | None = None       # None → kayıtlı anahtar korunur
    model: str
    tool_mode: str = "auto"          # auto | native | prompt
    extra_body: dict[str, Any] | None = None
    max_tokens: int | None = None    # yanıt token sınırı; boş = otomatik (bağlam penceresine göre, en çok 16384)
    vision: VisionIn = VisionIn()

    @field_validator("api_key")
    @classmethod
    def _k(cls, v): return _clean_key(v)

    @field_validator("base_url")
    @classmethod
    def _u(cls, v): return _clean_url(v)


class LlmTestIn(BaseModel):
    target: str = "main"             # main | vision
    llm: LlmIn


def _origin(url: str | None) -> str:
    """Anahtarın ait olduğu sunucu: şema + host + port (yol farkı önemsiz)."""
    from urllib.parse import urlsplit
    try:
        p = urlsplit((url or "").strip())
        return f"{p.scheme.lower()}://{(p.hostname or '').lower()}:{p.port or (443 if p.scheme == 'https' else 80)}"
    except ValueError:
        return ""


def _llm_keys(body: LlmIn) -> tuple[str, str]:
    """Formdan gelen anahtar yoksa kayıtlı (connections.json ya da .env) anahtar kullanılır — YALNIZ adres aynı sunucuyu
    gösteriyorsa. Adres başka bir sunucuya değiştiyse kayıtlı anahtar yeni sunucuya gönderilmez (anahtar sızmasın)."""
    cur = conns.llm_settings(get_settings())
    same_main = _origin(body.base_url) == _origin(cur.llm_base_url)
    main = body.api_key if body.api_key is not None else ((cur.llm_api_key or "") if same_main else "")
    if body.vision.api_key is not None:
        vis = body.vision.api_key
    else:
        cur_vis_url = cur.vision_base_url or cur.llm_base_url
        # görsel model ayrı anahtar tutmuyorsa ana modelin anahtarını kullanıyordur (aynı sunucu)
        cur_vis_key = cur.vision_api_key or ("" if cur.vision_base_url else cur.llm_api_key) or ""
        new_vis_url = body.base_url if body.vision.same_as_main else body.vision.base_url
        vis = cur_vis_key if _origin(new_vis_url) == _origin(cur_vis_url) else ""
    return main, vis


def _llm_target(body: LlmIn, target: str) -> tuple[str, str, str]:
    main_key, vis_key = _llm_keys(body)
    if target == "vision" and not body.vision.same_as_main:
        return body.vision.base_url.strip(), vis_key, body.vision.model.strip()
    return body.base_url.strip(), main_key, (body.vision.model if target == "vision" else body.model).strip()


@app.post("/api/settings/llm/models")
def llm_models(body: LlmTestIn, request: Request) -> dict[str, Any]:
    """Sunucudaki modeller (OpenAI uyumlu /models)."""
    _require_settings_access(request)
    from openai import OpenAI
    url, key, _ = _llm_target(body.llm, body.target)
    try:
        ms = OpenAI(base_url=url, api_key=key or "EMPTY", timeout=15, max_retries=0).models.list().data
        return {"ok": True, "models": sorted(m.id for m in ms)}
    except Exception as e:  # noqa: BLE001
        return {"ok": False, "error": _llm_error(e), "models": []}


_PIXEL = ("data:image/png;base64,iVBORw0KGgoAAAANSUhEUgAAAAgAAAAICAIAAABLbSncAAAAEklEQVR4nGP4z8CAFWEXHbQSACj/"
          "P8Hj9zQrAAAAAElFTkSuQmCC")


@app.post("/api/settings/llm/test")
def llm_test(body: LlmTestIn, request: Request) -> dict[str, Any]:
    """Modele gerçek kısa bir istek: ana model → metin + araç çağırma desteği; görsel model → küçük bir görsel."""
    _require_settings_access(request)
    import time as _t

    from openai import OpenAI
    url, key, model = _llm_target(body.llm, body.target)
    if not url or not model:
        return {"ok": False, "error": "API adresi ve model adı gerekli."}
    cli = OpenAI(base_url=url, api_key=key or "EMPTY", timeout=60, max_retries=0)
    extra = {"extra_body": body.llm.extra_body} if body.llm.extra_body else {}
    t0 = _t.perf_counter()
    try:
        if body.target == "vision":
            r = cli.chat.completions.create(model=model, max_tokens=40, temperature=0, messages=[{"role": "user", "content": [
                {"type": "text", "text": "Bu görselde hangi renk var? Tek kelimeyle yanıtla."},
                {"type": "image_url", "image_url": {"url": _PIXEL}}]}], **extra)
            return {"ok": True, "model": model, "elapsed_ms": int((_t.perf_counter() - t0) * 1000),
                    "reply": (r.choices[0].message.content or "").strip()[:120]}
        r = cli.chat.completions.create(model=model, max_tokens=20, temperature=0,
                                        messages=[{"role": "user", "content": "Yalnızca 'Tamam' yaz."}], **extra)
        out = {"ok": True, "model": model, "elapsed_ms": int((_t.perf_counter() - t0) * 1000),
               "reply": (r.choices[0].message.content or "").strip()[:120]}
    except Exception as e:  # noqa: BLE001
        return {"ok": False, "error": _llm_error(e)}
    try:  # araç çağırma (agent'ın veri / tasarım araçları için)
        tool = {"type": "function", "function": {"name": "kaydet", "description": "Bir sayıyı kaydeder",
                                                  "parameters": {"type": "object", "properties": {"sayi": {"type": "integer"}},
                                                                 "required": ["sayi"]}}}
        r = cli.chat.completions.create(model=model, max_tokens=200, temperature=0, tools=[tool],
                                        messages=[{"role": "user", "content": "kaydet aracını sayi=7 ile çağır."}], **extra)
        out["tools"] = "native" if r.choices[0].message.tool_calls else "prompt"
    except Exception:  # noqa: BLE001
        out["tools"] = "prompt"
    return out


def _llm_error(e: Exception) -> str:
    import openai
    if isinstance(e, openai.AuthenticationError):
        return "API anahtarı geçersiz ya da eksik (401)."
    if isinstance(e, openai.NotFoundError):
        return "Adres ya da model bulunamadı (404): API adresi genelde .../v1 ile biter; model adını kontrol edin."
    if isinstance(e, openai.APIConnectionError):
        why = type(e.__cause__).__name__ if e.__cause__ else ""
        if why == "LocalProtocolError":
            return "API anahtarında geçersiz karakter var (boşluk / satır sonu?). Anahtarı yeniden yapıştırın."
        return ("Sunucuya bağlanılamadı: adres doğru mu (http/https, port — ör. http://100.121.208.108:1234/v1), "
                "kurumsal proxy / sertifika engeli var mı?" + (f" [{why}]" if why else ""))
    if isinstance(e, UnicodeEncodeError):
        return "API anahtarında Türkçe ya da görünmez karakter var; anahtarı LM Studio'dan yeniden kopyalayın."
    if isinstance(e, openai.APITimeoutError):
        return "Sunucu zaman aşımına uğradı."
    if isinstance(e, openai.APIStatusError):
        return f"Sunucu hatası ({e.status_code}): {str(e)[:200]}"
    return f"{type(e).__name__}: {str(e)[:200]}"


@app.get("/api/settings/llm")
def get_llm(request: Request) -> dict[str, Any]:
    _require_settings_access(request)
    return conns.llm_effective(get_settings())


@app.put("/api/settings/llm")
def save_llm(body: LlmIn, request: Request) -> dict[str, Any]:
    ident = _require_settings_access(request)
    if not body.base_url.strip() or not body.model.strip():
        raise HTTPException(400, "API adresi ve model adı gerekli.")
    if body.tool_mode not in ("auto", "native", "prompt"):
        raise HTTPException(400, "Araç çağırma modu auto / native / prompt olmalı.")
    main_key, vis_key = _llm_keys(body)
    v = body.vision
    conns.save_connections({"llm": {
        "base_url": body.base_url.strip(), "model": body.model.strip(), "tool_mode": body.tool_mode,
        "extra_body": body.extra_body, "api_key_enc": conns.protect(main_key),
        "max_tokens": max(256, min(int(body.max_tokens), 131072)) if body.max_tokens else None,
        "vision": {"enabled": v.enabled and bool(v.model.strip()), "same_as_main": v.same_as_main, "model": v.model.strip(),
                   "base_url": "" if v.same_as_main else v.base_url.strip(),
                   "api_key_enc": "" if v.same_as_main else conns.protect(vis_key)},
    }})
    _rebuild_gateway()
    log.info("LLM bağlantısı güncellendi (%s): %s @ %s | görsel: %s", ident.username, body.model, body.base_url,
             v.model if v.enabled else "-")
    return {"ok": True, **state.gateway.health()}


@app.delete("/api/settings/llm")
def reset_llm(request: Request) -> dict[str, Any]:
    """Arayüzden kaydedilen LLM ayarını sil: .env'deki LLM_* / VISION_* ayarlarına dön."""
    _require_settings_access(request)
    conns.save_connections({"llm": None})
    _rebuild_gateway()
    return {"ok": True}


def _rebuild_gateway() -> None:
    state.gateway = LLMGateway(conns.llm_settings(get_settings()))
    if getattr(state, "agent", None) is not None:
        state.agent.llm = state.gateway


@app.get("/api/settings/connections/export")
def export_connections(request: Request) -> Response:
    """Geçerli bağlantı ayarlarını JSON olarak indirir (başka bir bilgisayara taşımak için).
    Şifreler dosyaya YAZILMAZ; içe aktarınca şifre yeniden girilir."""
    _require_settings_access(request)
    eff = conns.effective(get_settings())
    drop = {"has_password", "password", "password_enc"}
    clean = lambda d: {k: v for k, v in (d or {}).items() if k not in drop}  # noqa: E731
    dic = clean(eff["dictionary"])
    note = None
    if dic.get("kind") == "excel" and dic.get("excel_path"):
        note = ("Excel sözlük dosyası bu bilgisayardaki yolu gösterir; diğer bilgisayarda dosyayı yeniden yükleyin "
                "ya da ortak bir ağ yolu kullanın.")
    payload = {
        "app": "BI Lens", "type": "connection-settings", "version": 1,
        "exported_at": now_iso(), "exported_by": _me(request).username, "source": eff["source"],
        "data": clean(eff["data"]), "dictionary": dic,
        "llm": {k: v for k, v in conns.llm_effective(get_settings()).items() if k not in ("has_api_key", "source")} | {
            "vision": {k: v for k, v in conns.llm_effective(get_settings())["vision"].items() if k != "has_api_key"}},
        "notes": [n for n in ["Şifreler ve API anahtarları bu dosyada yoktur; içe aktardıktan sonra gerekenleri girin.",
                             note] if n],
    }
    body = json.dumps(payload, ensure_ascii=False, indent=2)
    return Response(body, media_type="application/json",
                    headers={"Content-Disposition": f'attachment; filename="bi_lens_baglanti_ayarlari_{now_iso()[:10]}.json"'})


@app.get("/api/settings/connections")
def get_connections(request: Request) -> dict[str, Any]:
    _require_settings_access(request)
    eff = conns.effective(get_settings())
    drivers = conns.installed_drivers()
    return {**eff, "drivers": [d for d in drivers if "SQL Server" in d],
            "driver": conns.best_sql_server_driver(drivers), "default_dictionary_db": conns.DEFAULT_DICTIONARY_DB,
            "startup_error": state.startup_error, "file": "backend/config/connections.json"}


@app.post("/api/settings/connections/test")
def test_connection(body: ConnTestIn, request: Request) -> dict[str, Any]:
    _require_settings_access(request)
    data, dic = _resolved(body)
    if body.target != "dictionary":
        try:
            info = _probe(data)
        except Exception as e:  # noqa: BLE001
            return {"ok": False, "error": friendly_error(str(e))}
        extras = [_probe_extra(data, db) for db in data.get("extra_databases") or []]
        bad = [x for x in extras if not x["ok"]]
        out = {"ok": not bad, **info, **({"extra_databases": extras} if extras else {})}
        if bad:
            out["error"] = "Ek veritabanlarına erişilemedi: " + "; ".join(f"{x['database']}: {x['error']}" for x in bad)
        return out
    # sözlük: kaynağa bağlan (SQL Server / MySQL / Excel), seçilen tablo / sayfaları gerçekten oku (rol başına satır sayısı)
    from app.dictionary.sources import collect
    if dic["kind"] == "none":
        return {"ok": True, "server_name": "—", "database": "Sözlük kullanılmıyor", "version": "yalnız veritabanı kataloğu",
                "dictionary_counts": {}, "warnings": ["Tablo ve view'lar veritabanından (yetkinize göre), açıklamalar MS_Description'dan, "
                                                      "ilişkiler foreign key'lerden gelir."]}
    try:
        reader, info = conns.open_dictionary_reader(dic)
    except Exception as e:  # noqa: BLE001
        return {"ok": False, "error": friendly_error(str(e))}
    catalog = _data_catalog(data)
    try:
        got = collect(reader, dic["sources"], dic.get("mappings"), catalog.get("known"))
    except Exception as e:  # noqa: BLE001
        return {"ok": False, **info, "error": "Bağlantı kuruldu ama sözlük okunamadı: " + friendly_error(str(e))}
    counts = got.counts
    info.update(dictionary_counts=counts, dictionary_tables=counts.get("tables", 0), warnings=got.warnings,
                derived_tables=got.derived_tables, sources_info=got.sources)
    if "relationships" not in got.explicit and got.rows.get("columns"):  # ilişki sözlüğü yok: otomatik kaç ilişki bulunur?
        try:
            n, src = DataDictionary(get_settings(), state.services.connector).preview_auto_relationships(
                got.rows.get("tables") or [], got.rows["columns"])
            counts["relationships"] = n
            info["relationship_source"] = src
        except Exception as e:  # noqa: BLE001
            log.info("Otomatik ilişki önizlemesi yapılamadı: %s", e)
    match = _dictionary_match(data, got, catalog)
    if match:
        info["catalog_match"] = match
    if got.errors:
        return {"ok": False, **info, "error": " ".join(got.errors)}
    return {"ok": True, **info}


def _probe_extra(data: dict[str, Any], database: str) -> dict[str, Any]:
    """Ek veritabanı: aynı sunucu / kimlik bilgisiyle bağlanılabiliyor mu, kaç nesneye SELECT yetkisi var."""
    from app.data.connector import SqlServerConnector
    try:
        con = SqlServerConnector(conns.build_odbc(data, database), 15)
        n = con.execute("SELECT COUNT(*) FROM sys.objects o WHERE o.type IN ('U', 'V') AND o.is_ms_shipped = 0 "
                        "AND HAS_PERMS_BY_NAME(QUOTENAME(SCHEMA_NAME(o.schema_id)) + '.' + QUOTENAME(o.name), 'OBJECT', 'SELECT') = 1",
                        1).rows[0][0]
        return {"database": database, "ok": True, "objects": int(n or 0)}
    except Exception as e:  # noqa: BLE001
        return {"database": database, "ok": False, "error": friendly_error(str(e))}


def _data_catalog(data: dict[str, Any]) -> dict[str, Any]:
    """Veri kaynağının kataloğu (nesneler + farklı kolon adları) — sözlük testi / kaydı için. Hata olursa {"error"}."""
    if not data.get("database"):
        return {}
    from app.data.connector import SqlServerConnector
    try:
        dd = DataDictionary(get_settings(), SqlServerConnector(data.get("odbc") or conns.build_odbc(data), 15))
        dd.extra_databases = list(data.get("extra_databases") or [])
        if not data.get("odbc"):   # formdaki bağlantı: ek veritabanları aynı kimlik bilgileriyle
            dd.extra_connector = lambda db: SqlServerConnector(conns.build_odbc(data, db), 15)
        cat = dd._read_catalog(columns=False)
        if cat is None:
            return {"error": dd.catalog_error}
        return {"objects": cat[0], "known": dd.known(cat, dd.column_names()), "extra_databases": dd.extra_databases}
    except Exception as e:  # noqa: BLE001
        return {"error": friendly_error(str(e))}


def _dictionary_match(data: dict[str, Any], got: Any, catalog: dict[str, Any]) -> dict[str, Any] | None:
    """Sözlükteki tablo / view adları veri kaynağında hangi nesnelere eşleniyor: bulunanlar, adı düzeltilenler,
    bulunamayanlar ve SELECT yetkisi olmayanlar (Ayarlar sayfasında gösterilir)."""
    names = sorted({str(r["table_name"]) for r in got.rows.get("tables") or [] if r.get("table_name")})
    if not names or not data.get("database"):
        return None
    from app.dictionary.names import NameResolver
    if "objects" not in catalog:
        return {"ok": False, "database": data["database"], "error": catalog.get("error")}
    objs = catalog["objects"]
    res = NameResolver(((s, o) for s, o, *_ in objs), catalog.get("extra_databases") or [])
    resolved = {res.resolve(n) for n in names}
    perms = {f"{s}.{o}".lower(): bool(p) for s, o, _t, p, *_ in objs}
    missing = sorted(res.missing.values())
    no_select = sorted(n for n in resolved if perms.get(n.lower()) is False)
    return {"ok": True, "database": data["database"], "total": len(resolved), "found": len(resolved) - len(missing),
            "missing": missing[:30], "missing_count": len(missing), "no_select": no_select[:30], "no_select_count": len(no_select),
            "renamed": [[a, b] for a, b in list(res.changed.items())[:10]], "renamed_count": len(res.changed),
            "db_objects": len(objs)}


@app.post("/api/settings/dictionary/tables")
def dictionary_candidate_tables(body: ConnTestIn, request: Request) -> dict[str, Any]:
    """Sözlük veritabanındaki tablolar + kolonlarına göre önerilen sözlük rolü."""
    _require_settings_access(request)
    _, dic = _resolved(body)
    from app.dictionary.sources import ROLES, _norm_key, suggest_role
    try:
        reader, _ = conns.open_dictionary_reader(dic)
        tables = reader.list_tables()
    except Exception as e:  # noqa: BLE001
        return {"ok": False, "error": friendly_error(str(e)), "tables": []}
    out = [{"name": n, "columns": cols, "role": suggest_role({_norm_key(c) for c in cols})} for n, cols in tables.items()]
    return {"ok": True, "tables": out,
            "roles": {k: {"label": v["label"], "required": v["required"], "optional": v["optional"], "must": v["must"]}
                      for k, v in ROLES.items()}}


@app.post("/api/settings/instances")
def find_instances(request: Request) -> dict[str, Any]:
    _require_settings_access(request)
    local = conns.local_instances()
    net = [x for x in conns.network_instances() if x.lower() not in {y.lower() for y in local}]
    return {"local": local, "network": net}


@app.post("/api/settings/databases")
def list_databases(body: ConnTestIn, request: Request) -> dict[str, Any]:
    _require_settings_access(request)
    data, dic = _resolved(body)
    fields = dic if body.target == "dictionary" else data
    if body.target == "dictionary" and dic["kind"] in ("mysql", "postgres"):
        try:
            from app.dictionary.sources import MySQLReader, PostgresReader
            r = (MySQLReader(dic["server"], dic.get("port") or 3306, dic["username"], dic.get("password") or "", "")
                 if dic["kind"] == "mysql" else
                 PostgresReader(dic["server"], dic.get("port") or 5432, dic["username"], dic.get("password") or "", "postgres",
                                ssl=bool(dic.get("encrypt"))))
            return {"ok": True, "databases": r.databases()}
        except Exception as e:  # noqa: BLE001
            return {"ok": False, "error": friendly_error(str(e)), "databases": []}
    try:
        from app.data.connector import SqlServerConnector
        r = SqlServerConnector(conns.build_odbc(fields, "master"), 15).execute(
            "SELECT name FROM sys.databases WHERE state = 0 AND HAS_DBACCESS(name) = 1 "
            "AND name NOT IN ('master','tempdb','model','msdb') ORDER BY name", 500)
        return {"ok": True, "databases": [row[0] for row in r.rows]}
    except Exception as e:  # noqa: BLE001
        return {"ok": False, "error": friendly_error(str(e)), "databases": []}


@app.put("/api/settings/connections")
def save_connections(body: ConnectionsIn, request: Request) -> dict[str, Any]:
    ident = _require_settings_access(request)
    data, dic = _resolved(body)
    if not data["database"]:
        raise HTTPException(400, "Veri kaynağı için veritabanı seçin.")
    # sözlük tabloları hatalıysa kaydetme: çalışan eski ayar korunsun (sunucuya o an ulaşılamıyorsa kayda izin ver)
    from app.dictionary import sources as dsrc
    try:
        got = None if dic["kind"] == "none" else dsrc.collect(conns.open_dictionary_reader(dic)[0], dic["sources"], dic.get("mappings"),
                                                               _data_catalog(data).get("known"))
    except Exception:  # noqa: BLE001
        got = None
    if got and got.errors:
        raise HTTPException(422, "Kaydedilmedi — sözlük tabloları hatalı: " + " ".join(got.errors))
    _remember_mappings(dic, got)
    store = lambda f: {k: v for k, v in f.items() if k != "password"} | {"password_enc": conns.protect(f.get("password") or "")}  # noqa: E731
    dic_section = ({"kind": "sqlserver", "same_as_data": True, "database": dic["database"], "sources": dic["sources"],
                    "mappings": dic["mappings"]}
                   if dic["kind"] == "sqlserver" and dic.get("same_as_data") else store(dic))
    conns.save_connections({"data": store(data), "dictionary": dic_section})  # llm bölümü korunur
    _rebuild_services()
    log.info("Bağlantı ayarları güncellendi (%s): veri=%s/%s sözlük=%s/%s", ident.username,
             data["server"], data["database"], dic["server"], dic["database"])
    return {"ok": state.startup_error is None, "error": state.startup_error,
            "tables": len(state.services.dictionary.tables)}


# ---------------------------------------------------------------- bölüm bazlı kayıt (veri kaynağı / sözlük ayrı ayrı)
class DataSectionIn(BaseModel):
    data: ConnIn


class DictSectionIn(BaseModel):
    dictionary: ConnIn
    data: ConnIn | None = None       # yalnız doğrulama için (kaydedilmez)


def _remember_mappings(dic: dict[str, Any], got: Any) -> None:
    """İçerikten bulunan sütun eşlemeleri de kaydedilir (elle seçilenlerle birlikte) — her açılışta yeniden aranmaz."""
    if not got:
        return
    merged = {k: dict(v) for k, v in (dic.get("mappings") or {}).items()}
    for name, m in got.learned_mappings().items():
        merged.setdefault(name, {}).update({f: h for f, h in m.items() if f not in merged.get(name, {})})
    dic["mappings"] = merged


def _store(f: dict[str, Any]) -> dict[str, Any]:
    return {k: v for k, v in f.items() if k not in ("password", "odbc")} | {"password_enc": conns.protect(f.get("password") or "")}


@app.put("/api/settings/connections/data")
def save_data_connection(body: DataSectionIn, request: Request) -> dict[str, Any]:
    """Yalnız veri kaynağını kaydeder (sözlük ve LLM ayarına dokunmaz)."""
    ident = _require_settings_access(request)
    data, _ = _resolved(ConnTestIn(target="data", data=body.data))
    if not data["database"]:
        raise HTTPException(400, "Veri kaynağı için veritabanı seçin.")
    conns.save_connections({"data": _store(data)})
    _rebuild_services()
    log.info("Veri kaynağı güncellendi (%s): %s/%s", ident.username, data["server"], data["database"])
    return {"ok": state.startup_error is None, "error": state.startup_error, "tables": len(state.services.dictionary.tables)}


@app.put("/api/settings/connections/dictionary")
def save_dictionary_connection(body: DictSectionIn, request: Request) -> dict[str, Any]:
    """Yalnız veri sözlüğünü kaydeder. "Veri kaynağıyla aynı" seçiliyse sunucu / kimlik bilgisi kopyalanmaz:
    çalışma anında veri kaynağının geçerli ayarı (arayüz ya da .env) kullanılır."""
    ident = _require_settings_access(request)
    saved = conns.load_connections() or {}
    _, dic = _resolved(ConnTestIn(target="dictionary", data=body.data or ConnIn(), dictionary=body.dictionary))
    if dic["kind"] == "sqlserver" and dic.get("same_as_data"):
        section = {"kind": "sqlserver", "same_as_data": True, "database": dic["database"], "sources": dic["sources"]}
    else:
        section = _store(dic)
    # içerikten tanıma için veri kaynağının kataloğu: formdaki veri kaynağı, yoksa kayıtlı / .env
    if body.data is not None:
        data_fields, _ = _resolved(ConnTestIn(target="data", data=body.data))
    else:
        data_fields = {"odbc": conns.data_odbc(get_settings()), "database": "(geçerli veri kaynağı)"}
    # sözlük tabloları hatalıysa kaydetme (sunucuya o an ulaşılamıyorsa kayda izin ver); doğrulama çalışma anındaki bağlantıyla
    from app.dictionary import sources as dsrc
    runtime = conns.dictionary_conn_fields({"data": saved.get("data"), "dictionary": section})
    runtime["kind"] = dic["kind"]
    if dic["kind"] != "sqlserver" or not dic.get("same_as_data"):
        runtime = dic
    try:
        got = None if dic["kind"] == "none" else dsrc.collect(conns.open_dictionary_reader(runtime)[0], dic["sources"],
                                                               dic.get("mappings"), _data_catalog(data_fields).get("known"))
    except Exception:  # noqa: BLE001
        got = None
    if got and got.errors:
        raise HTTPException(422, "Kaydedilmedi — sözlük tabloları hatalı: " + " ".join(got.errors))
    _remember_mappings(dic, got)
    section["mappings"] = dic["mappings"]
    conns.save_connections({"dictionary": section})
    _rebuild_services()
    log.info("Veri sözlüğü güncellendi (%s): %s %s", ident.username, dic["kind"], dic.get("database", ""))
    return {"ok": state.startup_error is None, "error": state.startup_error, "tables": len(state.services.dictionary.tables)}


@app.delete("/api/settings/connections/{section}")
def reset_connection_section(section: str, request: Request) -> dict[str, Any]:
    """Tek bölümü sil (data → .env SQLSERVER_ODBC, dictionary → dictionary.toml)."""
    _require_settings_access(request)
    if section not in ("data", "dictionary"):
        raise HTTPException(404, "Bölüm data ya da dictionary olmalı.")
    conns.save_connections({section: None})
    _rebuild_services()
    return {"ok": state.startup_error is None, "error": state.startup_error}


@app.delete("/api/settings/connections")
def reset_connections(request: Request) -> dict[str, Any]:
    """Arayüzden kaydedilen veritabanı ayarlarını sil: .env / dictionary.toml'a geri dön (LLM ayarı korunur)."""
    _require_settings_access(request)
    conns.save_connections({"data": None, "dictionary": None})
    _rebuild_services()
    return {"ok": state.startup_error is None, "error": state.startup_error}


def _rebuild_services() -> None:
    state.services = build_services()
    state.agent.services = state.services


@app.exception_handler(ValidationError)
def _validation_handler(_: Request, exc: ValidationError) -> JSONResponse:
    return JSONResponse({"detail": [f"{'.'.join(map(str, e['loc']))}: {e['msg']}" for e in exc.errors()]}, status_code=422)


# --------------------------------------------------------------------------- Vitrin: yayınlama, paylaşım, izleme
class _Snapshot:
    """Yayınlanmış sürüm: dashboard verisi fonksiyonları tasarım oturumu gibi kullanır (spec + dataset'ler + rol).
    Rol, raporu AÇAN kullanıcının veri rolüdür (yayınlayanınki değil)."""

    def __init__(self, spec: ReportSpec, datasets: list[Dataset], role: str, model_relationships: list | None = None):
        self.spec, self.datasets, self.user_role = spec, datasets, role
        self.model_relationships = model_relationships or []


def _snapshot(version: dict[str, Any], role: str) -> _Snapshot:
    return _Snapshot(ReportSpec.model_validate(version["spec"]),
                     [Dataset.model_validate(d) for d in version["datasets"]], role, version.get("model_relationships"))


def _public_spec(spec: ReportSpec) -> dict[str, Any]:
    """İzleyiciye giden spec: SQL yok (dataset'lerden yalnız alan tanımları)."""
    out = spec.model_dump(exclude_none=True)
    out["datasets"] = [{"id": d.id, "description": d.description, "sql": "", "fields": [f.model_dump(exclude_none=True)
                                                                                      for f in d.fields]}
                       for d in spec.datasets]
    return out


def _report_card(rep: dict[str, Any], ident: Identity, grants: list[dict[str, Any]]) -> dict[str, Any]:
    return {"id": rep["report_id"], "title": rep["title"], "description": rep.get("description"),
            "owner": rep["owner"], "owner_name": rep.get("owner_name") or rep["owner"], "version": rep["current_version"],
            "status": rep["status"], "domains": rep.get("domains") or [], "created_at": rep["created_at"],
            "updated_at": rep["updated_at"], "mine": authz.owns(ident, rep["owner"]),
            "can_export": authz.can_export_report(ident, rep, grants),
            "can_manage": (manage := authz.at_least(ident, "admin") or authz.owns(ident, rep["owner"])),
            "shared_with": len(grants) if manage else None,
            "session_id": rep["session_id"] if manage and authz.at_least(ident, "builder") else None}   # "Tasarımda aç"


def _viewable(request: Request, report_id: str) -> tuple[dict[str, Any], Identity, list[dict[str, Any]]]:
    ident = _me(request)
    meta = _meta()
    rep = meta.report(report_id)
    grants = meta.grants(report_id) if rep else []
    if not rep or not authz.can_view_report(ident, rep, grants):
        raise HTTPException(404, "Rapor bulunamadı ya da görüntüleme yetkiniz yok.")
    return rep, ident, grants


def _manageable(request: Request, report_id: str) -> tuple[dict[str, Any], Identity]:
    ident = _me(request)
    rep = _meta().report(report_id)
    if not rep:
        raise HTTPException(404, "Rapor bulunamadı.")
    if not (authz.at_least(ident, "admin") or authz.owns(ident, rep["owner"])):
        raise HTTPException(403, "Bu yayını yalnızca sahibi ya da yönetici değiştirebilir.")
    return rep, ident


class GrantIn(BaseModel):
    principal_type: str              # user | group
    principal: str
    can_export: bool = False

    @field_validator("principal_type")
    @classmethod
    def _t(cls, v: str) -> str:
        if v not in ("user", "group"):
            raise ValueError("principal_type user ya da group olmalı")
        return v

    @field_validator("principal")
    @classmethod
    def _p(cls, v: str) -> str:
        v = " ".join(v.split())
        if not v or len(v) > 256:
            raise ValueError("kullanıcı / grup adı boş olamaz (en çok 256 karakter)")
        return v


class PublishIn(BaseModel):
    description: str | None = None
    notes: str | None = None             # sürüm notu
    grants: list[GrantIn] | None = None  # None → mevcut paylaşım korunur


@app.get("/api/sessions/{sid}/publication")
def session_publication(sid: str, request: Request) -> dict[str, Any]:
    """Tasarım oturumunun Vitrin durumu: yayında mı, hangi sürüm, kimlerle paylaşıldı, yayında olmayan değişiklik var mı."""
    s, ident = _own_session(request, sid)
    meta = _meta()
    rep = meta.report_by_session(sid)
    if not rep:
        return {"published": False}
    cur = meta.version(rep["report_id"])
    changed = bool(s.spec) and cur is not None and (
        cur["spec"] != s.spec.model_dump(mode="json")
        or cur["datasets"] != [d.model_dump(mode="json") for d in s.datasets]
        or (cur.get("model_relationships") or []) != s.model_relationships)
    return {"published": True, "report": _report_card(rep, ident, meta.grants(rep["report_id"])),
            "grants": meta.grants(rep["report_id"]), "versions": meta.versions(rep["report_id"]), "unpublished_changes": changed}


@app.post("/api/sessions/{sid}/publish")
def publish_session(sid: str, body: PublishIn, request: Request) -> dict[str, Any]:
    """Tasarımın o anki halini Vitrin'e yeni sürüm olarak yayınlar (spec + dataset'ler değişmez kopya)."""
    s, ident = _own_session(request, sid)
    meta = _meta()
    if not s.spec:
        raise HTTPException(409, "Yayınlamak için önce dashboard oluşturun.")
    kinds = state.services.dictionary._table_kinds()
    from app.harness.session import _source_tables
    domains = _report_domains(_source_tables([d.model_dump() for d in s.datasets]), kinds)
    prev = meta.report_by_session(sid)
    rep = meta.publish(session_id=sid, owner=(prev or {}).get("owner") or s.owner or ident.username,
                       owner_name=(prev or {}).get("owner_name") or s.owner_name or ident.display_name,
                       title=s.title, description=(body.description if body.description is not None
                                                   else (prev or {}).get("description")),
                       domains=domains, spec=s.spec.model_dump(mode="json"),
                       datasets=[d.model_dump(mode="json") for d in s.datasets], notes=body.notes, by=ident.username,
                       model_relationships=s.model_relationships)
    if body.grants is not None:
        meta.set_grants(rep["report_id"], [g.model_dump() for g in body.grants], ident.username)
    with state.store.lock(sid):
        s = _session(sid)
        s.status = "live"
        s.add(TranscriptItem(role="system", content=f"Vitrin'de yayınlandı: sürüm {rep['current_version']} ({ident.display_name})"))
        state.store.save(s)
    _audit(ident, "report_publish", rep["report_id"], version=rep["current_version"], session=sid,
           grants=None if body.grants is None else len(body.grants))
    return {"ok": True, "report": _report_card(rep, ident, meta.grants(rep["report_id"])), "session": s.public()}


@app.post("/api/vitrin/{report_id}/retire")
def retire_report(report_id: str, request: Request) -> dict[str, Any]:
    rep, ident = _manageable(request, report_id)
    _meta().set_status(report_id, "retired")
    try:
        with state.store.lock(rep["session_id"]):
            s = state.store.get(rep["session_id"])
            if s.status == "live":
                s.status = "test"
                s.add(TranscriptItem(role="system", content=f"Vitrin'den kaldırıldı ({ident.display_name})"))
                state.store.save(s)
    except KeyError:
        pass
    _audit(ident, "report_retire", report_id)
    return {"ok": True}


@app.get("/api/vitrin/{report_id}/grants")
def get_grants(report_id: str, request: Request) -> list[dict[str, Any]]:
    _manageable(request, report_id)
    return _meta().grants(report_id)


@app.put("/api/vitrin/{report_id}/grants")
def put_grants(report_id: str, body: list[GrantIn], request: Request) -> list[dict[str, Any]]:
    _rep, ident = _manageable(request, report_id)
    meta = _meta()
    before = meta.grants(report_id)
    meta.set_grants(report_id, [g.model_dump() for g in body], ident.username)
    after = meta.grants(report_id)
    _audit(ident, "grants_change", report_id, before=[f"{g['principal_type']}:{g['principal']}" for g in before],
           after=[f"{g['principal_type']}:{g['principal']}" for g in after])
    return after


class OwnerIn(BaseModel):
    owner: str
    owner_name: str | None = None


@app.put("/api/vitrin/{report_id}/owner")
def transfer_owner(report_id: str, body: OwnerIn, request: Request) -> dict[str, Any]:
    """Sahiplik devri (yalnız admin): yayın + tasarım oturumu yeni sahibe geçer."""
    ident = _need(request, "admin")
    meta = _meta()
    rep = meta.report(report_id)
    if not rep:
        raise HTTPException(404, "Rapor bulunamadı.")
    owner = " ".join(body.owner.split())
    if not owner:
        raise HTTPException(400, "Yeni sahip boş olamaz.")
    meta.set_owner(report_id, owner, body.owner_name or owner)
    try:
        with state.store.lock(rep["session_id"]):
            s = state.store.get(rep["session_id"])
            s.owner, s.owner_name = owner, body.owner_name or owner
            state.store.save(s)
    except KeyError:
        pass
    _audit(ident, "owner_transfer", report_id, old=rep["owner"], new=owner)
    return {"ok": True}


@app.get("/api/vitrin")
def vitrin_list(request: Request) -> list[dict[str, Any]]:
    """Vitrin: kullanıcının görebildiği yayınlar (izin verilen, kendisinin ya da admin için hepsi)."""
    ident = _me(request)
    meta = _meta()
    grants = meta.all_grants()
    out = []
    for rep in meta.reports():
        g = grants.get(rep["report_id"], [])
        if authz.can_view_report(ident, rep, g):
            out.append(_report_card(rep, ident, g))
    return out


@app.get("/api/vitrin/{report_id}")
def vitrin_report(report_id: str, request: Request) -> dict[str, Any]:
    """Yayınlanmış rapor: spec (SQL'siz), filtreler, veri tarihi — izleyicinin veri rolüyle."""
    rep, ident, grants = _viewable(request, report_id)
    ver = _meta().version(report_id)
    if not ver:
        raise HTTPException(404, "Yayın sürümü bulunamadı.")
    snap = _snapshot(ver, ident.role)
    bindings = _bindings(snap)
    defs = _filter_defs(snap, bindings)
    for f in defs:
        f["options"] = _filter_options(f["key"], ident.role) if f["key"] else []
    try:
        data_date = _data_date_info(snap, ident.role)
    except Exception as e:  # noqa: BLE001
        log.warning("Veri tarihi aralığı alınamadı: %s", e)
        data_date = None
    _audit(ident, "report_view", report_id, version=ver["version"])
    return {"report": _report_card(rep, ident, grants), "spec": _public_spec(snap.spec),
            "version": {"version": ver["version"], "notes": ver.get("notes"), "published_by": ver["published_by"],
                        "published_at": ver["published_at"]},
            "filters": {"filters": defs, "bindings": bindings, **({"data_date": data_date} if data_date else {})}}


@app.post("/api/vitrin/{report_id}/data")
def vitrin_data(report_id: str, body: DashboardQuery, request: Request) -> dict[str, Any]:
    _rep, ident, _g = _viewable(request, report_id)
    ver = _meta().version(report_id)
    if not ver:
        raise HTTPException(404, "Yayın sürümü bulunamadı.")
    return _dashboard_data(_snapshot(ver, ident.role), [sel.model_dump() for sel in body.selections])


@app.get("/api/vitrin/{report_id}/export/html")
def vitrin_export(report_id: str, request: Request) -> HTMLResponse:
    rep, ident, grants = _viewable(request, report_id)
    if not authz.can_export_report(ident, rep, grants):
        raise HTTPException(403, "Bu raporu dışa aktarma izniniz yok.")
    ver = _meta().version(report_id)
    if not ver:
        raise HTTPException(404, "Yayın sürümü bulunamadı.")
    viewer = get_settings().viewer_html
    if not viewer.exists():
        raise HTTPException(503, "Viewer derlenmemiş: frontend klasöründe `npm run build:viewer` çalıştırın.")
    snap = _snapshot(ver, ident.role)
    data = _dashboard_data(snap)
    data["bindings"] = _bindings(snap)
    data["filter_keys"] = {f["id"]: f["key"] for f in _filter_defs(snap, data["bindings"])}
    payload = json.dumps({"spec": _public_spec(snap.spec), "data": data},
                         ensure_ascii=False, default=str).replace("<", "\\u003c")
    html = viewer.read_text(encoding="utf-8").replace("__REPORT_JSON__", payload, 1)
    _audit(ident, "report_export", report_id, version=ver["version"])
    title = rep["title"] or "dashboard"
    ascii_name = "".join(c if c.isascii() and c.isalnum() else "_" for c in title)[:60] or "dashboard"
    return HTMLResponse(html, headers={"Content-Disposition": f"attachment; filename=\"{ascii_name}.html\"; "
                                                              f"filename*=UTF-8''{quote(title[:60])}.html"})


# --------------------------------------------------------------------------- yönetim (admin)
class AssignmentIn(BaseModel):
    principal_type: str
    principal: str
    platform_role: str | None = None
    data_role: str | None = None

    @field_validator("principal_type")
    @classmethod
    def _t(cls, v: str) -> str:
        if v not in ("user", "group"):
            raise ValueError("principal_type user ya da group olmalı")
        return v


@app.get("/api/admin/overview")
def admin_overview(request: Request) -> dict[str, Any]:
    _need(request, "admin")
    meta = _meta()
    settings = get_settings()
    grants = meta.all_grants()
    return {"mode": settings.platform_mode, "store": meta.kind,
            "platform_admins": [a.strip() for a in settings.platform_admins.split(",") if a.strip()],
            "platform_roles": list(authz.RANK), "data_roles": sorted(state.services.policies),
            "assignments": [a.__dict__ for a in meta.assignments()],
            "reports": [{**r, "grants": grants.get(r["report_id"], [])} for r in meta.reports()]}


@app.put("/api/admin/assignments")
def put_assignment(body: AssignmentIn, request: Request) -> dict[str, Any]:
    ident = _need(request, "admin")
    if body.platform_role is not None and body.platform_role not in authz.RANK:
        raise HTTPException(400, f"Platform rolü {list(authz.RANK)} olmalı.")
    if body.data_role is not None and body.data_role not in state.services.policies:
        raise HTTPException(400, f"Veri rolü policy.toml'da tanımlı olmalı: {sorted(state.services.policies)}")
    principal = " ".join(body.principal.split())
    if not principal:
        raise HTTPException(400, "Kullanıcı / grup adı boş olamaz.")
    _meta().set_assignment(body.principal_type, principal, body.platform_role, body.data_role, ident.username)
    _audit(ident, "role_assign", None, principal_type=body.principal_type, principal=principal,
           platform_role=body.platform_role, data_role=body.data_role)
    return {"ok": True}


@app.delete("/api/admin/assignments/{principal_type}/{principal:path}")
def delete_assignment(principal_type: str, principal: str, request: Request) -> dict[str, Any]:
    ident = _need(request, "admin")
    ok = _meta().delete_assignment(principal_type, principal)
    if ok:
        _audit(ident, "role_unassign", None, principal_type=principal_type, principal=principal)
    return {"ok": ok}


@app.get("/api/admin/audit")
def admin_audit(request: Request, limit: int = 200, user: str | None = None, event: str | None = None,
                report: str | None = None) -> list[dict[str, Any]]:
    _need(request, "admin")
    return _meta().audit_events(limit, user or None, event or None, report or None)


# --------------------------------------------------------------------------- platform veritabanı ayarı
class MetaConnIn(BaseModel):
    meta: ConnIn


@app.get("/api/settings/platform")
def get_platform_settings(request: Request) -> dict[str, Any]:
    _require_settings_access(request)
    settings = get_settings()
    return {"mode": settings.platform_mode, "store": state.meta.kind if state.meta else None, "error": state.meta_error,
            "meta": conns.meta_public(settings)}


@app.put("/api/settings/platform")
def save_platform_settings(body: MetaConnIn, request: Request) -> dict[str, Any]:
    """Platform veritabanını kaydeder: önce bağlanıp tabloları hazırlar (yoksa oluşturur), başarılıysa kaydeder."""
    ident = _require_settings_access(request)
    saved = (conns.load_connections() or {}).get("meta")
    f = _fields(body.meta, saved)
    if not f["database"]:
        raise HTTPException(400, "Platform veritabanı adını girin (ör. BI_Lens_Meta).")
    try:
        st = MetaStore.sqlserver(conns.build_odbc(f))
    except Exception as e:  # noqa: BLE001
        return {"ok": False, "error": friendly_error(str(e))[:400]}
    conns.save_connections({"meta": _store(f)})
    state.meta, state.meta_error = st, None
    log.info("Platform veritabanı güncellendi (%s): %s/%s", ident.username, f["server"], f["database"])
    _audit(ident, "settings_platform_db", None, server=f["server"], database=f["database"])
    return {"ok": True, "error": None}


@app.delete("/api/settings/platform")
def reset_platform_settings(request: Request) -> dict[str, Any]:
    _require_settings_access(request)
    conns.save_connections({"meta": None})
    state.meta = build_meta()
    return {"ok": state.meta is not None, "error": state.meta_error}


# --------------------------------------------------------------------------- arayüz (derlenmiş)
# frontend/dist varsa backend arayüzü de sunar: Node.js / Vite olmadan tek port (kurum / çevrimdışı kurulum).
# Geliştirmede Vite (5173) kullanılmaya devam eder; bu yalnızca http://127.0.0.1:<port>/ adresini de çalışır kılar.
_DIST = BACKEND_DIR.parent / "frontend" / "dist"
if (_DIST / "index.html").exists():
    from fastapi.staticfiles import StaticFiles

    class _Frontend(StaticFiles):
        """index.html önbelleğe alınmaz (yeni sürümde eski arayüz gelmesin); assets/ dosyaları adında hash taşır."""

        async def get_response(self, path, scope):  # type: ignore[override]
            resp = await super().get_response(path, scope)
            if resp.media_type == "text/html" or path in ("", ".", "index.html"):
                resp.headers["Cache-Control"] = "no-cache"
            return resp

    app.mount("/", _Frontend(directory=_DIST, html=True), name="frontend")
